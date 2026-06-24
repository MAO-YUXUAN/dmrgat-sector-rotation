from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
GAT_METRICS = ROOT / "data" / "processed" / "metrics.json"
CNN_METRICS = ROOT / "CNN" / "cnn_metrics.json"
LSTM_METRICS = ROOT / "LSTM" / "lstm_metrics.json"
OUTPUT_PATH = ROOT / "DMRGAT_CNN_LSTM_Comparison_Section.docx"


def load_json(path: Path) -> dict:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def set_style(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)


def add_para(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)


def fmt(value: float | None) -> str:
    if value is None:
        return "N/A"
    return f"{value:.4f}"


def metric(metrics: dict, key: str) -> float | None:
    value = metrics.get(key)
    if value is None:
        return None
    return float(value)


def add_metrics_table(document: Document, gat: dict, cnn: dict, lstm: dict) -> None:
    rows = [
        ("DMRGAT", gat),
        ("CNN", cnn),
        ("LSTM", lstm),
    ]
    table = document.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = ["Model", "Train RMSE", "Val RMSE", "Test RMSE", "Test MAE", "Test IC", "Test RankIC"]
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)
                run.bold = True

    for model_name, metrics in rows:
        cells = table.add_row().cells
        values = [
            model_name,
            fmt(metric(metrics, "train_rmse")),
            fmt(metric(metrics, "val_rmse")),
            fmt(metric(metrics, "test_rmse")),
            fmt(metric(metrics, "test_mae")),
            fmt(metric(metrics, "test_ic")),
            fmt(metric(metrics, "test_rank_ic")),
        ]
        for cell, value in zip(cells, values):
            cell.text = value
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)


def main() -> None:
    gat = load_json(GAT_METRICS)
    cnn = load_json(CNN_METRICS)
    lstm = load_json(LSTM_METRICS)

    doc = Document()
    set_style(doc)
    doc.add_heading("5.3 Comparison with CNN and LSTM Baselines", level=1)

    add_para(
        doc,
        "To evaluate whether the proposed dynamic multi-relation graph attention architecture provides additional predictive value, this study compares DMRGAT with two non-graph neural-network baselines: a convolutional neural network (CNN) and a long short-term memory network (LSTM). The comparison is placed in the empirical simulation section because it concerns out-of-sample forecasting performance rather than the mathematical definition of the proposed algorithm. All three models are trained and evaluated under the same sector-rotation task: predicting each industry's future 30-trading-day excess return over the CSI 300 benchmark.",
    )
    add_para(
        doc,
        "The experimental design is kept consistent across the three models. DMRGAT, CNN, and LSTM use the same Shenwan level-1 industry universe, the same 15-dimensional node feature vector, the same target variable, and the same chronological train-validation-test split. This consistency is important because the purpose of the comparison is to isolate the effect of model architecture rather than differences in data preprocessing or sample selection. In particular, the DMRGAT model uses the dynamic multi-relation graph and edge-weighted attention mechanism, the CNN model treats each daily industry-feature matrix as a two-dimensional input, and the LSTM model uses historical sequences of industry features to capture temporal dependence.",
    )
    add_para(
        doc,
        "The three baselines therefore represent different modeling assumptions. DMRGAT assumes that sector performance is affected by both node-level characteristics and inter-industry relationships. CNN assumes that useful local patterns may exist in the industry-feature matrix, but it does not explicitly use graph edges or economic relation priors. LSTM assumes that industry performance is mainly driven by temporal patterns in historical features. Comparing these models helps clarify whether graph-based relational modeling is more suitable than purely convolutional or recurrent neural architectures for the present sector-rotation problem.",
    )

    doc.add_heading("Table 1. Forecasting Performance of DMRGAT, CNN, and LSTM", level=2)
    add_metrics_table(doc, gat, cnn, lstm)

    add_para(
        doc,
        "The empirical results show that the CNN baseline achieves the lowest test RMSE and test MAE among the three models. This indicates that CNN provides the best point-forecast accuracy in terms of numerical error. However, its test RankIC is negative, which means that the ordering of predicted industry returns is not aligned with the realized cross-sectional ordering in the test period. For sector-rotation applications, this distinction is important because investment decisions usually depend more on ranking industries correctly than on minimizing average point-prediction error.",
    )
    add_para(
        doc,
        "The LSTM baseline obtains a relatively high validation RankIC, but its test RMSE and MAE are substantially larger than those of CNN and DMRGAT, and its test RankIC decreases to a value close to zero. This pattern suggests that the recurrent model captures some temporal structure in the validation period but does not generalize strongly to the final test period. Such behavior is common in financial forecasting, where temporal patterns may be regime-dependent and unstable across market environments. The LSTM result therefore highlights the difficulty of relying only on historical sequential information for sector rotation.",
    )
    add_para(
        doc,
        "DMRGAT does not produce the lowest point-prediction error, but it achieves a slightly positive test RankIC. This result is consistent with the design objective of the model. DMRGAT is not intended merely to minimize numerical prediction error; it is designed to learn relative sector strength through graph-based information propagation and a ranking-oriented loss function. The positive test RankIC indicates that the model retains a weak but directionally meaningful cross-sectional ranking signal in the out-of-sample period. Although the magnitude of this signal remains modest, it is more aligned with the practical objective of sector allocation than a negative ranking correlation.",
    )
    add_para(
        doc,
        "Overall, the comparison suggests that different neural architectures emphasize different aspects of the forecasting task. CNN is effective at reducing point-forecast error, LSTM attempts to capture temporal dependence but exhibits weaker out-of-sample stability, and DMRGAT provides the most relevant ranking interpretation for sector rotation. From an investment perspective, the DMRGAT result is meaningful because sector selection is primarily a cross-sectional ranking problem. Nevertheless, the weak magnitude of the test RankIC also indicates that the current graph-learning framework should be interpreted as an exploratory modeling approach rather than a mature trading signal.",
    )
    add_para(
        doc,
        "These findings support the methodological motivation of introducing graph structure into industry-level forecasting. Sector returns are influenced not only by each industry's own historical features but also by changing dependencies among industries. By incorporating return co-movement, capital-flow similarity, and static economic linkage into the attention mechanism, DMRGAT provides an interpretable way to model these dependencies. The comparison with CNN and LSTM therefore strengthens the empirical discussion: even when non-graph models perform well in terms of error metrics, graph-based modeling remains valuable when the research objective is ranking-based sector rotation.",
    )

    doc.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
