from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "DMRGAT_Revised_Introduction.docx"


def set_style(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)


def add_para(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)


def add_bullet(document: Document, text: str) -> None:
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)


def main() -> None:
    doc = Document()
    set_style(doc)
    doc.add_heading("1. Introduction", level=1)

    add_para(
        doc,
        "Sector rotation forecasting is an important problem in quantitative investment and tactical asset allocation. In real markets, investors rarely evaluate industries in isolation. Instead, they compare the relative attractiveness of different sectors as macroeconomic expectations, liquidity conditions, policy signals, and market sentiment evolve. When growth expectations improve, cyclical sectors may respond earlier and outperform the market; when risk aversion rises, defensive sectors may become relatively more resilient. Therefore, the objective of sector rotation is not simply to predict whether a single industry will rise or fall. A more relevant task is to identify which industries are likely to outperform others over a given investment horizon.",
    )
    add_para(
        doc,
        "This relative nature makes sector rotation a cross-sectional forecasting problem. Traditional approaches often rely on each industry's own historical return, volatility, momentum, valuation, or trading-volume information. These variables are useful, but they do not fully capture the relational structure of financial markets. Industries are connected through production chains, shared macroeconomic exposure, capital-flow migration, investor attention, and policy transmission. For example, a change in commodity prices may affect upstream resource sectors and downstream manufacturing industries simultaneously; technology-sector sentiment may propagate across electronics, computers, and communications; and financial conditions may influence banks, non-bank financial institutions, real estate, and cyclical industries through credit channels. Ignoring such inter-industry dependence may limit the ability of a model to understand how sector strength is formed and transmitted.",
    )
    add_para(
        doc,
        "Recent machine-learning methods provide flexible tools for financial forecasting, but many implementations still treat industries as independent samples or rely on fixed feature matrices without explicitly modeling economic relations. Graph neural networks offer a natural way to represent inter-sector dependence, because each industry can be treated as a node and relations between industries can be represented as edges. However, a key challenge remains: the relationship between industries is not single-dimensional. Some links arise from short-term return co-movement, some from synchronized trading activity or capital-flow behavior, and others from relatively stable industrial or supply-chain connections. A single static graph may therefore be insufficient for sector rotation, where market relations are dynamic and economically heterogeneous.",
    )
    add_para(
        doc,
        "To address this issue, this paper proposes a Dynamic Multi-Relation Graph Attention Network (DMRGAT) for industry-level sector rotation forecasting. The model represents each industry as a node and constructs a dynamic graph by combining three relation channels: rolling return correlation, rolling similarity in trading-amount changes as a proxy for capital-flow synchronization, and a static industry-prior graph based on economic linkage. The fused graph is then used in an edge-weighted graph attention mechanism, where the attention score between two industries is guided by both learned node representations and the corresponding graph edge weight. This design allows the model to learn adaptive information aggregation while preserving economically meaningful relation priors.",
    )
    add_para(
        doc,
        "Unlike a pure classification framework, the present study formulates the prediction target as future excess return relative to the CSI 300 benchmark. This formulation is more consistent with practical allocation decisions because portfolio managers usually need to rank industries by expected relative strength rather than merely classify them into outperforming and underperforming groups. Accordingly, the model is trained with a ranking-oriented objective that emphasizes pairwise ordering across industries while retaining a small regression component to stabilize the prediction scale. Evaluation is conducted using both point-error metrics and cross-sectional ranking metrics, especially RankIC.",
    )
    add_para(
        doc,
        "The study is implemented on 15 core Shenwan level-1 industries in the Chinese A-share market. The empirical design uses a 30-trading-day forecasting horizon and a chronological train-validation-test split. In addition to the proposed DMRGAT model, convolutional neural network (CNN) and long short-term memory (LSTM) baselines are constructed under the same feature set, target definition, and sample split. This comparison is used not to claim that the graph model dominates every metric, but to examine whether graph-based relational modeling provides a useful and interpretable framework for ranking-oriented sector allocation.",
    )
    add_para(
        doc,
        "This paper makes the following contributions:",
    )
    add_bullet(
        doc,
        "First, it formulates sector rotation as a dynamic graph-based cross-sectional forecasting problem rather than a collection of independent industry-level predictions.",
    )
    add_bullet(
        doc,
        "Second, it constructs a multi-relation industry graph that integrates return co-movement, capital-flow similarity, and static economic linkage, thereby reflecting different channels of inter-industry dependence.",
    )
    add_bullet(
        doc,
        "Third, it introduces an edge-weighted graph attention mechanism in which fused relation weights guide the attention scores, improving the economic interpretability of information aggregation.",
    )
    add_bullet(
        doc,
        "Fourth, it aligns the learning objective with the investment purpose of sector rotation by using a ranking-oriented loss and evaluating the model with RankIC in addition to RMSE and MAE.",
    )
    add_bullet(
        doc,
        "Fifth, it provides an empirical comparison with CNN and LSTM baselines and uses attention heatmaps to examine the relational patterns learned by the graph model.",
    )
    add_para(
        doc,
        "The empirical results should be interpreted with appropriate caution. Financial excess-return prediction is a low signal-to-noise task, and the current out-of-sample ranking signal remains modest. The main value of the proposed framework is therefore methodological and interpretive: it provides a coherent way to connect dynamic market linkage, graph attention, ranking-oriented learning, and sector-allocation decision support. This positioning is important because a useful investment research model does not only need to generate forecasts; it should also offer a transparent structure for understanding how those forecasts are formed.",
    )
    add_para(
        doc,
        "The remainder of this paper is organized as follows. Section 2 presents the basic model and problem setup, including the feature construction, target definition, and dynamic graph formulation. Section 3 introduces the DMRGAT algorithm and explains the edge-weighted attention mechanism. Section 4 provides theoretical analysis of graph fusion, ranking consistency, and optimization behavior. Section 5 reports the simulation results and compares DMRGAT with CNN and LSTM baselines. Section 6 discusses the practical application of the model as a sector-rotation decision-support tool. Section 7 concludes the paper and outlines directions for future research.",
    )

    doc.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
