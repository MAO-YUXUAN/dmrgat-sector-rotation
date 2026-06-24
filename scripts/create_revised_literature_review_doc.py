from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "DMRGAT_Revised_Literature_Review.docx"


def set_style(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)


def add_para(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)


def main() -> None:
    doc = Document()
    set_style(doc)
    doc.add_heading("2. Literature Review", level=1)

    add_para(
        doc,
        "Financial return prediction has long been a central topic in empirical asset pricing, quantitative investment, and machine learning. Early forecasting studies relied heavily on linear time-series models, such as autoregressive models and ARIMA-type specifications. These models are useful for describing simple temporal dependence, but they are often insufficient for financial markets because return dynamics are nonlinear, noisy, and regime-dependent. With the development of deep learning, recurrent neural networks, convolutional neural networks, attention mechanisms, and graph neural networks have been introduced into financial forecasting. The following review summarizes the main streams of related research and clarifies the gap addressed by the proposed DMRGAT framework.",
    )

    doc.add_heading("2.1 Financial Time-Series Forecasting with LSTM and Transformer Models", level=2)
    add_para(
        doc,
        "Recurrent neural networks were introduced to capture temporal dependence in sequential data, but standard RNNs suffer from vanishing and exploding gradients when modeling long-term dependencies. Long Short-Term Memory (LSTM), proposed by Hochreiter and Schmidhuber, addresses this limitation through gating mechanisms and memory cells. Because financial prices and returns are naturally sequential, LSTM has been widely applied to stock price prediction, index forecasting, volatility modeling, and trend classification. Many empirical studies report that LSTM-based models perform competitively against traditional linear models and shallow machine-learning baselines, particularly when the target depends on nonlinear temporal patterns.",
    )
    add_para(
        doc,
        "Subsequent research has extended LSTM models in several directions. Some studies improve the input representation by incorporating technical indicators, investor sentiment, market microstructure variables, or dimensionality-reduction methods such as principal component analysis. Others combine LSTM with CNN, attention modules, random forests, particle swarm optimization, or other hybrid structures to capture both local feature patterns and long-range temporal dependence. These extensions reflect an important insight: the predictive performance of recurrent models depends not only on sequence modeling ability but also on feature quality and model architecture.",
    )
    add_para(
        doc,
        "Transformer-based models provide another route for financial time-series prediction. Unlike recurrent models, Transformers use self-attention to model long-range dependencies and support parallel training. Recent studies have applied Transformer variants to stock prediction, multi-scale feature learning, market relation modeling, and reinforcement-learning-based portfolio decisions. However, both LSTM and Transformer models are primarily designed for temporal dependence. When applied directly to sector or stock prediction, they may fail to explicitly represent cross-sectional dependence among assets. For sector rotation, this limitation is important because industries are connected through supply chains, shared macro exposures, capital flows, and investor attention. Therefore, sequence models are useful baselines but do not fully address the relational nature of industry-level allocation.",
    )

    doc.add_heading("2.2 Graph Neural Networks for Financial Market Prediction", level=2)
    add_para(
        doc,
        "Graph neural networks provide a natural framework for modeling financial markets as relational systems. In a financial graph, nodes may represent stocks, firms, industries, or time intervals, while edges may represent ownership relations, supply-chain links, industry membership, return correlations, news co-occurrence, or transaction connections. Graph convolutional networks, GraphSAGE, gated graph neural networks, and related architectures have been applied to stock ranking, price movement prediction, risk contagion analysis, and portfolio selection. Compared with independent time-series models, GNNs can aggregate information from related nodes and capture cross-sectional information transmission.",
    )
    add_para(
        doc,
        "Existing GNN-based financial studies have shown that relational information can improve market prediction. For example, corporate relation networks can help models learn firm representations; industry or supply-chain graphs can improve stock ranking; and knowledge graphs can incorporate textual or fundamental information into market forecasting. These studies demonstrate that financial assets should not be treated as isolated units. Market shocks, capital flows, and investor expectations often propagate through relational structures, and graph models provide a flexible way to encode such propagation.",
    )
    add_para(
        doc,
        "Nevertheless, basic GNN models have several limitations. First, many models rely on static graphs defined by prior knowledge, such as industry classifications or fixed corporate relations. Static graphs may not reflect the evolving nature of financial markets, where correlations, liquidity linkages, and investor attention change over time. Second, some GNNs use fixed aggregation rules, meaning that all connected neighbors contribute according to predetermined weights or normalized graph structure. This may be too restrictive for sector rotation, where the importance of a neighboring industry can vary across market regimes. Third, many studies focus on stock-level prediction rather than industry-level rotation, leaving sector allocation relatively underexplored from a dynamic graph-learning perspective.",
    )

    doc.add_heading("2.3 Graph Attention and Multi-Relation Financial Graphs", level=2)
    add_para(
        doc,
        "Graph Attention Networks (GATs) address part of the fixed-aggregation limitation by learning attention weights over neighboring nodes. Instead of treating all neighbors equally, GAT assigns adaptive importance scores based on node representations. This makes GAT particularly attractive for financial markets because the influence between assets is asymmetric, time-varying, and heterogeneous. Financial applications of GAT and attention-based graph models include stock return prediction, volume forecasting, relation-aware stock ranking, event propagation, semantic company networks, and dynamic transaction networks.",
    )
    add_para(
        doc,
        "A further development is the use of multi-relation or heterogeneous graphs. Financial relationships rarely come from a single source. Two companies or industries may be related because they belong to the same sector, share supply-chain exposure, move together in returns, receive similar capital flows, or appear together in news. Multi-relation graph models attempt to preserve these different channels instead of collapsing them into one relation type. Some studies aggregate multiple graphs, learn relation-specific attention, or combine graph embeddings with temporal modules such as LSTM. These approaches provide useful evidence that relation diversity matters in financial prediction.",
    )
    add_para(
        doc,
        "However, several gaps remain. Many attention-based financial graph models still rely on static or semi-static graph structures. Others use multi-relation information but combine it through simple concatenation or unstructured weighted summation without explicitly linking the weights to economic interpretation. In addition, attention weights are often treated only as post-hoc visualization tools, while the graph construction itself may not impose meaningful economic priors. For sector rotation, a desirable model should combine dynamic market-based relations, static industry-prior relations, and adaptive attention in a unified framework. This is the motivation for the dynamic multi-relation graph design used in DMRGAT.",
    )

    doc.add_heading("2.4 Ranking-Oriented Sector Allocation", level=2)
    add_para(
        doc,
        "A separate but important issue concerns the learning objective. Many financial prediction studies optimize point-wise loss functions, such as mean squared error, mean absolute error, or classification accuracy. These metrics are appropriate when the goal is to predict the exact value or direction of a single asset. However, portfolio allocation and sector rotation are often ranking-oriented tasks. Investors need to decide which industries should be overweighted or underweighted relative to a benchmark. Therefore, the relative ordering of predicted returns may be more important than the absolute numerical forecast.",
    )
    add_para(
        doc,
        "In quantitative investment, information coefficient and rank information coefficient are commonly used to evaluate the cross-sectional association between model signals and future returns. A model with a low point-forecast error may still be unsuitable for allocation if it ranks future winners below future losers. Conversely, a model with modest numerical accuracy may still be useful if it generates a stable positive ranking signal. This distinction is especially relevant for sector rotation because the number of industries is limited and the decision is usually based on relative strength. Ranking losses and RankIC-based evaluation therefore provide a more investment-aligned framework than simple classification accuracy.",
    )
    add_para(
        doc,
        "Despite this practical importance, many deep-learning studies in financial prediction still emphasize classification accuracy or regression error. Ranking-oriented objectives are less frequently integrated with graph attention mechanisms, particularly in industry-level forecasting. The present study addresses this gap by predicting future benchmark-relative excess returns and training the model with a hybrid objective dominated by pairwise ranking loss. This design aligns the model with the economic purpose of sector allocation rather than treating prediction as a purely statistical curve-fitting problem.",
    )

    doc.add_heading("2.5 Research Gap and Positioning of This Study", level=2)
    add_para(
        doc,
        "The literature suggests three major gaps. First, sequence models such as LSTM and Transformer are effective for temporal dependence but do not explicitly model inter-industry relations. Second, basic GNNs introduce graph structure but often rely on static graphs or fixed aggregation mechanisms. Third, GAT-based financial models learn adaptive neighbor importance but frequently lack dynamic multi-relation graph construction and ranking-oriented optimization. These gaps are particularly relevant for sector rotation, where industries interact through multiple economic channels and the investment decision depends on cross-sectional ranking.",
    )
    add_para(
        doc,
        "This paper contributes to the literature by proposing a Dynamic Multi-Relation Graph Attention Network for sector rotation forecasting. The model constructs a fused industry graph from rolling return correlation, trading-amount similarity as a capital-flow proxy, and static industry-prior linkage. The fused edge weight is introduced directly into the attention score, allowing the model to combine learned node representations with economically meaningful relation priors. The prediction target is future excess return over the CSI 300 benchmark, and the training objective emphasizes cross-sectional ranking through pairwise ranking loss.",
    )
    add_para(
        doc,
        "Relative to LSTM and Transformer models, DMRGAT explicitly represents cross-industry dependence. Relative to basic GNNs, it uses adaptive attention rather than fixed neighbor aggregation. Relative to vanilla GAT, it incorporates dynamic multi-relation edge weights and links attention learning to economically interpretable graph construction. Relative to many financial prediction studies, it evaluates the model under a ranking-oriented framework that better reflects sector-allocation decisions. In this sense, the proposed framework is positioned as a finance-oriented graph-learning approach for interpretable, ranking-based industry rotation research.",
    )

    doc.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
