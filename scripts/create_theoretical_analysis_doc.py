from __future__ import annotations

import json
from html import escape
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
CONFIG_PATH = ROOT / "configs" / "default.json"
OUTPUT_PATH = ROOT / "DMRGAT_Theoretical_Analysis_Section.docx"


def load_json(path: Path) -> dict:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def mr(text: str) -> str:
    return f"<m:r><m:t>{escape(text)}</m:t></m:r>"


def sub(base: str, lower: str) -> str:
    return f"<m:sSub><m:e>{base}</m:e><m:sub>{lower}</m:sub></m:sSub>"


def sup(base: str, upper: str) -> str:
    return f"<m:sSup><m:e>{base}</m:e><m:sup>{upper}</m:sup></m:sSup>"


def subsup(base: str, lower: str, upper: str) -> str:
    return f"<m:sSubSup><m:e>{base}</m:e><m:sub>{lower}</m:sub><m:sup>{upper}</m:sup></m:sSubSup>"


def frac(num: str, den: str) -> str:
    return f"<m:f><m:num>{num}</m:num><m:den>{den}</m:den></m:f>"


def omath(expr: str) -> str:
    return f'<m:oMathPara {nsdecls("m")}><m:oMath>{expr}</m:oMath></m:oMathPara>'


def add_equation(document: Document, expr: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p._p.append(parse_xml(omath(expr)))


def add_para(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)


def add_mixed_para(document: Document, segments: list[str | tuple[str, str]]) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    for segment in segments:
        if isinstance(segment, tuple):
            kind, expr = segment
            if kind != "math":
                raise ValueError(f"Unknown segment kind: {kind}")
            p._p.append(parse_xml(f'<m:oMath {nsdecls("m")}>{expr}</m:oMath>'))
        else:
            r = p.add_run(segment)
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)


def set_style(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)


def eq_fusion(alpha: float, beta: float, gamma: float) -> str:
    return (
        sub(mr("A"), mr("t"))
        + mr("=")
        + mr(f"{alpha:.2f}")
        + sub(mr("A_corr"), mr("t"))
        + mr("+")
        + mr(f"{beta:.2f}")
        + mr("A_industry")
        + mr("+")
        + mr(f"{gamma:.2f}")
        + sub(mr("A_fund"), mr("t"))
    )


def eq_convex(alpha: float, beta: float, gamma: float) -> str:
    return mr(f"{alpha:.2f}+{beta:.2f}+{gamma:.2f}=1") + mr(",  ") + mr("alpha,beta,gamma>=0")


def eq_bound() -> str:
    return mr("0<=") + sub(mr("A"), mr("ij,t")) + mr("<=1")


def eq_attention_bound() -> str:
    eij = sub(mr("e"), mr("ij,t"))
    sij = sub(mr("s"), mr("ij,t"))
    aij = sub(mr("A"), mr("ij,t"))
    return mr("|") + eij + mr("|=|") + aij + sij + mr("|<=|") + sij + mr("|")


def eq_rank_loss() -> str:
    yi = sub(mr("y"), mr("i"))
    yj = sub(mr("y"), mr("j"))
    yhi = sub(mr("ŷ"), mr("i"))
    yhj = sub(mr("ŷ"), mr("j"))
    spread = mr("-(") + yi + mr("-") + yj + mr(")(") + yhi + mr("-") + yhj + mr(")")
    return mr("ell_ij=log(1+exp(") + spread + mr("))")


def eq_rank_monotonic() -> str:
    yi = sub(mr("y"), mr("i"))
    yj = sub(mr("y"), mr("j"))
    yhi = sub(mr("ŷ"), mr("i"))
    yhj = sub(mr("ŷ"), mr("j"))
    return yi + mr(">") + yj + mr(" and ") + yhi + mr(">") + yhj


def eq_total_loss() -> str:
    return mr("L=4.0") + sub(mr("L"), mr("rank")) + mr("+0.05") + sub(mr("L"), mr("mse"))


def eq_update() -> str:
    return sub(mr("theta"), mr("k+1")) + mr("=") + sub(mr("theta"), mr("k")) + mr("-eta") + sub(mr("nabla"), mr("theta")) + mr("L(") + sub(mr("theta"), mr("k")) + mr(")")


def eq_stationary() -> str:
    return mr("||") + sub(mr("nabla"), mr("theta")) + mr("L(") + sup(mr("theta"), mr("*")) + mr(")|| approx 0")


def main() -> None:
    config = load_json(CONFIG_PATH)
    alpha = float(config.get("alpha", 0.45))
    beta = float(config.get("beta", 0.25))
    gamma = float(config.get("gamma", 0.30))

    doc = Document()
    set_style(doc)
    doc.add_heading("4. Theoretical Analysis", level=1)

    add_para(
        doc,
        "For a journal-style manuscript, the theoretical analysis should not simply repeat the computational steps of the algorithm. Its purpose is to clarify why the proposed modeling structure is mathematically coherent, economically interpretable, and suitable for the sector-rotation objective. In this study, the theoretical discussion focuses on four aspects: the convex fusion of multiple relation matrices, the stability and interpretation of edge-weighted attention, the consistency between the ranking loss and the investment objective, and the practical convergence properties of the optimization procedure. These arguments do not constitute a global optimality proof, because the full neural network remains non-convex. Instead, they provide a principled justification for the architecture and explain why the implemented model is a reasonable graph-learning framework for industry-level relative return prediction.",
    )

    doc.add_heading("4.1 Convex Fusion and Graph Stability", level=2)
    add_mixed_para(
        doc,
        [
            "The fused adjacency matrix is constructed as a convex combination of three economically distinct relation channels. In the current implementation, the fusion rule is",
        ],
    )
    add_equation(doc, eq_fusion(alpha, beta, gamma))
    add_equation(doc, eq_convex(alpha, beta, gamma))
    add_mixed_para(
        doc,
        [
            "This condition is theoretically useful because it keeps the final graph within the feasible set spanned by the three base relation matrices. The return-correlation graph captures recent price co-movement, the fund-flow graph captures synchronized changes in trading-amount dynamics, and the industry-prior graph captures relatively stable economic linkage. A convex fusion does not create an unrestricted artificial relation matrix; rather, it reallocates relative emphasis among interpretable information channels. Therefore, each entry of ",
            ("math", sub(mr("A"), mr("t"))),
            " can still be interpreted as an edge-strength prior formed from observable market relations and structural economic knowledge.",
        ],
    )
    add_mixed_para(
        doc,
        [
            "After the correlation matrices are transformed into nonnegative edge-strength matrices, the implemented fused graph satisfies the boundedness condition",
        ],
    )
    add_equation(doc, eq_bound())
    add_para(
        doc,
        "The boundedness property is important for numerical stability. If graph weights were allowed to take large or negative values, the attention score could be excessively amplified or sign-distorted before softmax normalization. Such behavior would make training more sensitive to initialization and learning rate. By contrast, bounded nonnegative edge weights allow the graph to act as a prior strength matrix rather than as an uncontrolled scaling factor. This is especially relevant in financial applications because rolling correlations and liquidity proxies may be noisy. The convex formulation limits the influence of any single noisy relation source and makes the fused graph more comparable across trading dates.",
    )

    doc.add_heading("4.2 Edge-Weighted Attention as Prior-Guided Learning", level=2)
    add_mixed_para(
        doc,
        [
            "The DMRGAT layer differs from a vanilla GAT because the raw neural attention score is multiplied by the fused edge weight. Let ",
            ("math", sub(mr("s"), mr("ij,t"))),
            " denote the feature-driven attention score before graph weighting. The implemented edge-weighted score can be written as",
        ],
    )
    add_equation(
        doc,
        sub(mr("e"), mr("ij,t"))
        + mr("=")
        + sub(mr("A"), mr("ij,t"))
        + sub(mr("s"), mr("ij,t")),
    )
    add_para(
        doc,
        "where the score is subsequently normalized over the neighborhood of each node. This formulation can be interpreted as prior-guided attention learning. The neural component learns from node features, while the graph component constrains the attention mechanism toward economically meaningful neighbor relations. In other words, the model does not search for attention weights in a completely unconstrained space. It learns adaptive attention under a relation prior determined by return co-movement, capital-flow similarity, and industry linkage.",
    )
    add_equation(doc, eq_attention_bound())
    add_para(
        doc,
        "The inequality shows that when the fused edge weight lies between zero and one, the graph prior cannot increase the absolute magnitude of the raw feature-driven attention score. It can only preserve or attenuate that score. This does not guarantee superior prediction, because an incorrect prior may still suppress useful information. However, it provides a stabilizing mechanism: weak or economically unsupported edges are less likely to dominate attention purely because of random feature noise. Compared with a standard GAT, DMRGAT therefore introduces a structured inductive bias. Compared with a fixed GNN aggregation rule, it remains adaptive because the final attention coefficients are still learned from node representations and normalized locally.",
    )

    doc.add_heading("4.3 Ranking Consistency and Investment Interpretation", level=2)
    add_para(
        doc,
        "The prediction target is the future excess return of each industry over the benchmark, but the practical sector-rotation decision is primarily a ranking problem. Investors usually care about whether the model can place stronger industries above weaker industries, rather than whether it can perfectly estimate the numerical value of every future return. The training objective is therefore designed to align with this economic interpretation.",
    )
    add_equation(doc, eq_rank_loss())
    add_mixed_para(
        doc,
        [
            "This pairwise loss decreases when the predicted ordering is consistent with the realized ordering. For example, if the realized excess return of industry ",
            ("math", mr("i")),
            " is higher than that of industry ",
            ("math", mr("j")),
            ", the ideal ranking relation is",
        ],
    )
    add_equation(doc, eq_rank_monotonic())
    add_para(
        doc,
        "When this ordering holds, the product between the realized spread and the predicted spread is positive, causing the loss term to become smaller. When the predicted ordering is reversed, the product becomes negative and the loss increases. This property directly connects the training criterion to RankIC evaluation, because both emphasize cross-sectional ordering rather than absolute return matching. The ranking loss is therefore theoretically consistent with the purpose of sector allocation: it encourages the model to produce a useful relative-strength signal.",
    )
    add_equation(doc, eq_total_loss())
    add_para(
        doc,
        "The small MSE component plays a regularization role. A pure ranking loss can identify ordering but may leave the scale of predictions weakly controlled. Adding a small regression term stabilizes the predicted return magnitude while keeping the main objective focused on relative ranking. This hybrid objective is appropriate for empirical finance, where the level of future returns is noisy but the relative ordering of sectors can still contain economically useful information.",
    )

    doc.add_heading("4.4 Optimization and Local Convergence", level=2)
    add_para(
        doc,
        "The full DMRGAT model is nonlinear because it combines learned projections, attention normalization, nonlinear activations, dropout, and a regression head. Therefore, the overall loss is non-convex, and a global convergence guarantee is not available. This limitation is shared by most deep learning models, including CNN, LSTM, and vanilla GAT baselines. The relevant theoretical question is not whether the optimizer can find the global optimum, but whether the objective is differentiable and whether gradient-based optimization can approach a locally stable stationary point under standard practical conditions.",
    )
    add_equation(doc, eq_update())
    add_mixed_para(
        doc,
        [
            "The ranking component is smooth because it is based on the log-exp function, and the MSE component is quadratic. The attention operation is differentiable almost everywhere with respect to model parameters. Under bounded input features, finite graph size, bounded edge weights, and a sufficiently controlled learning rate, gradient-based optimization can be expected to reduce the empirical objective until it reaches an approximate stationary point:",
        ],
    )
    add_equation(doc, eq_stationary())
    add_para(
        doc,
        "In implementation, Adam and early stopping are used to improve practical convergence behavior. Early stopping is important because financial samples are noisy and regime-dependent; continuing to reduce training loss may weaken out-of-sample ranking performance. The validation RankIC is therefore used as the main model-selection criterion. This choice is coherent with the ranking-based objective and avoids selecting a model solely because it achieves lower in-sample numerical error.",
    )

    doc.add_heading("4.5 Theoretical Comparison with CNN, LSTM, and Vanilla GAT", level=2)
    add_para(
        doc,
        "The theoretical role of DMRGAT can be further clarified by comparing it with alternative neural architectures. A CNN baseline can learn local patterns in the industry-feature matrix, but it does not explicitly represent economic links among industries. Its convolutional filters depend on the arbitrary ordering of industries in the input matrix, and the model has no built-in mechanism to distinguish whether two sectors are economically connected. An LSTM baseline can capture temporal dependence in industry features, but it treats each industry's historical sequence as the central source of information and does not directly model cross-industry propagation. A vanilla GAT introduces graph attention, but without edge-weighted priors it learns attention mainly from node representations and may allocate attention to relations that are statistically convenient but economically less interpretable.",
    )
    add_para(
        doc,
        "DMRGAT occupies an intermediate and finance-oriented position. It is more structured than CNN and LSTM because it explicitly encodes inter-industry dependence. It is also more economically constrained than vanilla GAT because the attention score is guided by a fused relation prior. This design may reduce flexibility relative to a completely unconstrained attention mechanism, but it improves interpretability and imposes a useful inductive bias in a low-signal financial prediction environment. The theoretical value of the model is therefore not that it guarantees higher accuracy in every market regime, but that it aligns the architecture, graph construction, and training objective with the economic nature of sector rotation.",
    )
    add_para(
        doc,
        "In summary, the theoretical analysis supports the proposed framework from four perspectives. Convex graph fusion keeps relation weights interpretable and numerically bounded. Edge-weighted attention introduces a prior-guided mechanism for information aggregation. Pairwise ranking loss aligns the optimization target with cross-sectional sector allocation. Finally, the convergence discussion clarifies that the model should be evaluated under local optimization and out-of-sample ranking stability rather than unrealistic global optimality claims. These properties provide a rigorous foundation for the empirical comparison reported in the simulation section.",
    )

    doc.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
