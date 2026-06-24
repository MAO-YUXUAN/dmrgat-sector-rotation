from __future__ import annotations

from html import escape
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "DMRGAT_Section_3_2_Algorithm_OMML_Equations.docx"


def mr(text: str) -> str:
    return f"<m:r><m:t>{escape(text)}</m:t></m:r>"


def sub(base: str, lower: str) -> str:
    return f"<m:sSub><m:e>{base}</m:e><m:sub>{lower}</m:sub></m:sSub>"


def sup(base: str, upper: str) -> str:
    return f"<m:sSup><m:e>{base}</m:e><m:sup>{upper}</m:sup></m:sSup>"


def subsup(base: str, lower: str, upper: str) -> str:
    return f"<m:sSubSup><m:e>{base}</m:e><m:sub>{lower}</m:sub><m:sup>{upper}</m:sup></m:sSubSup>"


def frac(num: str, den: str) -> str:
    return f"<m:f><m:num>{num}</m:num><m:den>{den}</m:den></m:f>"


def nary_sum(lower: str, expr: str) -> str:
    return sub(mr("∑"), lower) + expr


def neighbor_i() -> str:
    return mr("j∈") + sub(mr("N"), mr("i"))


def k_neighbor_i() -> str:
    return mr("k∈") + sub(mr("N"), mr("i"))


def omath(expr: str) -> str:
    return f'<m:oMathPara {nsdecls("m")}><m:oMath>{expr}</m:oMath></m:oMathPara>'


def add_equation(document: Document, expr: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p._p.append(parse_xml(omath(expr)))


def add_para(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)


def add_mixed_para(document: Document, segments: list[str | tuple[str, str]]) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    for segment in segments:
        if isinstance(segment, tuple):
            kind, expr = segment
            if kind != "math":
                raise ValueError(f"Unknown segment kind: {kind}")
            p._p.append(parse_xml(f'<m:oMath {nsdecls("m")}>{expr}</m:oMath>'))
        else:
            r = p.add_run(segment)
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)


def set_style(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)


def eq_h0() -> str:
    return subsup(mr("h"), mr("i"), mr("(0)")) + mr("=") + sub(mr("x"), mr("i,t"))


def eq_z() -> str:
    return sub(mr("z"), mr("i")) + mr("=") + mr("W") + subsup(mr("h"), mr("i"), mr("(0)"))


def eq_e() -> str:
    left = sub(mr("e"), mr("ij,t"))
    aij = sub(mr("A"), mr("ij,t"))
    zi = sub(mr("z"), mr("i"))
    zj = sub(mr("z"), mr("j"))
    attention = (
        mr("LeakyReLU")
        + mr("(")
        + sup(mr("a"), mr("T"))
        + mr("[")
        + zi
        + mr(" ∥ ")
        + zj
        + mr("]")
        + mr(")")
    )
    return left + mr("=") + aij + mr("·") + attention


def eq_alpha() -> str:
    left = sub(mr("α"), mr("ij,t"))
    num = mr("exp") + mr("(") + sub(mr("e"), mr("ij,t")) + mr(")")
    den_term = mr("exp") + mr("(") + sub(mr("e"), mr("ik,t")) + mr(")")
    den = nary_sum(k_neighbor_i(), den_term)
    return left + mr("=") + frac(num, den)


def eq_h1() -> str:
    left = subsup(mr("h"), mr("i"), mr("(1)"))
    lower = neighbor_i()
    alpha = sup(sub(mr("α"), mr("ij,t")), mr("(m)"))
    zj = subsup(mr("z"), mr("j"), mr("(m)"))
    inner_sum = nary_sum(lower, alpha + zj)
    inner = mr("σ") + mr("(") + inner_sum + mr(")")
    concat = subsup(mr("Concat"), mr("m=1"), mr("4")) + mr("[") + inner + mr("]")
    return left + mr("=") + concat


def eq_h2() -> str:
    left = subsup(mr("h"), mr("i"), mr("(2)"))
    lower = neighbor_i()
    alpha = sup(sub(mr("α"), mr("ij,t")), mr("(2)"))
    w2 = sup(mr("W"), mr("(2)"))
    h1 = subsup(mr("h"), mr("j"), mr("(1)"))
    inner_sum = nary_sum(lower, alpha + w2 + h1)
    return left + mr("=") + mr("σ") + mr("(") + inner_sum + mr(")")


def eq_yhat() -> str:
    yhat = subsup(mr("ŷ"), mr("i,t"), mr("(30)"))
    return yhat + mr("=") + sub(mr("W"), mr("o")) + subsup(mr("h"), mr("i"), mr("(2)"))


def eq_lrank() -> str:
    left = sub(mr("L"), mr("rank"))
    mean = sub(mr("mean"), mr("i,j"))
    yi = sub(mr("y"), mr("i"))
    yj = sub(mr("y"), mr("j"))
    yhi = sub(mr("ŷ"), mr("i"))
    yhj = sub(mr("ŷ"), mr("j"))
    spread = mr("-(") + yi + mr("-") + yj + mr(")(") + yhi + mr("-") + yhj + mr(")")
    rhs = mean + mr(" log(1+exp(") + spread + mr("))")
    return left + mr("=") + rhs


def eq_lmse() -> str:
    left = sub(mr("L"), mr("mse"))
    term = mr("(") + sub(mr("ŷ"), mr("i")) + mr("-") + sub(mr("y"), mr("i")) + sup(mr(")"), mr("2"))
    rhs = frac(mr("1"), mr("N")) + nary_sum(mr("i"), term)
    return left + mr("=") + rhs


def eq_ltotal() -> str:
    return mr("L=4.0") + sub(mr("L"), mr("rank")) + mr("+0.05") + sub(mr("L"), mr("mse"))


def math_symbol(text: str) -> str:
    return mr(text)


def main() -> None:
    doc = Document()
    set_style(doc)
    doc.add_heading("3.2 DMRGAT Architecture and Algorithm", level=1)

    add_mixed_para(
        doc,
        [
            "The proposed DMRGAT model is designed to transform the dynamic multi-relation industry graph into a set of predictive node representations. Unlike a conventional time-series model that treats each industry independently, DMRGAT explicitly incorporates cross-industry dependence through graph attention. At each trading date ",
            ("math", mr("t")),
            ", the model takes the node feature matrix ",
            ("math", sub(mr("X"), mr("t"))),
            " and the fused adjacency matrix ",
            ("math", sub(mr("A"), mr("t"))),
            " as inputs. Each row of ",
            ("math", sub(mr("X"), mr("t"))),
            " corresponds to one industry, and each element of ",
            ("math", sub(mr("A"), mr("t"))),
            " represents the strength of the relation between two industries after integrating return correlation, capital-flow similarity, and static industry linkage. The purpose of the model is to learn a nonlinear mapping from the dynamic graph state ",
            ("math", mr("(") + sub(mr("G"), mr("t")) + mr(", ") + sub(mr("X"), mr("t")) + mr(")")),
            " to the predicted future excess return for each industry node.",
        ],
    )
    add_mixed_para(
        doc,
        [
            "The initial node representation of industry ",
            ("math", mr("i")),
            " is defined as follows. This notation indicates that the input of the first graph attention layer is exactly the feature vector constructed in the data preprocessing stage.",
        ],
    )
    add_equation(doc, eq_h0())
    add_mixed_para(
        doc,
        [
            "Before computing attention scores, the raw node representation ",
            ("math", subsup(mr("h"), mr("i"), mr("(0)"))),
            " is projected into a hidden feature space through a trainable linear transformation:",
        ],
    )
    add_equation(doc, eq_z())
    add_mixed_para(
        doc,
        [
            "This projection maps heterogeneous financial variables, including returns, volatility, momentum, liquidity-related information, and cross-sectional rankings, into a common latent representation. It also allows the model to learn task-specific combinations of these variables rather than relying directly on their raw scales. In the current implementation, the hidden dimension is set to ",
            ("math", mr("32")),
            ", which is appropriate for a relatively small industry graph with ",
            ("math", mr("15")),
            " nodes.",
        ],
    )
    add_mixed_para(
        doc,
        [
            "For each edge ",
            ("math", mr("(") + mr("i") + mr(",") + mr("j") + mr(")")),
            ", the weighted attention score is computed by combining neural attention with the prior graph relation:",
        ],
    )
    add_equation(doc, eq_e())
    add_mixed_para(
        doc,
        [
            "This equation is the key difference between DMRGAT and a vanilla GAT. In a standard GAT, the attention score is determined only by the transformed node features. In DMRGAT, the score is further multiplied by the fused edge weight ",
            ("math", sub(mr("A"), mr("ij,t"))),
            ". Therefore, the model does not allocate attention in a completely unconstrained manner. Instead, attention allocation is guided by economically meaningful prior relations. Industries with stronger return co-movement, capital-flow similarity, or structural linkage receive greater prior influence in the attention mechanism.",
        ],
    )
    add_mixed_para(
        doc,
        [
            "The attention score is then normalized over the neighborhood ",
            ("math", sub(mr("N"), mr("i"))),
            " of node ",
            ("math", mr("i")),
            " through a softmax operation:",
        ],
    )
    add_equation(doc, eq_alpha())
    add_mixed_para(
        doc,
        [
            "The normalized coefficient ",
            ("math", sub(mr("α"), mr("ij,t"))),
            " can be interpreted as the relative importance of industry ",
            ("math", mr("j")),
            " when updating the representation of industry ",
            ("math", mr("i")),
            ". Since the coefficients are normalized within the neighborhood, the attention mechanism produces a probability-like distribution over neighboring industries. This makes the model adaptive to local graph structure and allows different industries to depend on different sets of influential neighbors.",
        ],
    )
    add_mixed_para(
        doc,
        [
            "Using the normalized attention weights, the first-layer node representation ",
            ("math", subsup(mr("h"), mr("i"), mr("(1)"))),
            " is computed through neighborhood aggregation. In the current implementation, the first graph attention layer uses ",
            ("math", mr("M=4")),
            " attention heads:",
        ],
    )
    add_equation(doc, eq_h1())
    add_mixed_para(
        doc,
        [
            "The multi-head design is meaningful in the sector-rotation setting because industry relationships are heterogeneous. One head may learn cyclical co-movement, another may capture technology-sector spillovers, while another may focus on financial or consumption-related linkages. Instead of forcing all relations into a single attention distribution, the model allows different heads to represent different relational perspectives. The head-specific attention coefficients ",
            ("math", sup(sub(mr("α"), mr("ij,t")), mr("(m)"))),
            " generated by the implementation provide an interpretable view of these different relational patterns.",
        ],
    )
    add_mixed_para(
        doc,
        [
            "After the first graph attention layer, the concatenated representation ",
            ("math", subsup(mr("h"), mr("i"), mr("(1)"))),
            " is normalized, passed through a nonlinear activation, and regularized by dropout. The resulting representation is then passed into a second graph attention layer with one attention head. Conceptually, the first layer aggregates information from directly connected neighboring industries, while the second layer further integrates the information already propagated through the first layer. Therefore, a two-layer GAT structure allows each industry node to incorporate both first-order neighborhood information and higher-order relational signals transmitted through neighboring industries.",
        ],
    )
    add_equation(doc, eq_h2())
    add_mixed_para(
        doc,
        [
            "Here, the superscript in ",
            ("math", subsup(mr("h"), mr("i"), mr("(1)"))),
            " and ",
            ("math", subsup(mr("h"), mr("i"), mr("(2)"))),
            " denotes the graph attention layer index rather than the number of training iterations. This distinction is important because the model depth describes representation transformations during one forward pass, while optimization iterations refer to repeated parameter updates during training.",
        ],
    )
    add_mixed_para(
        doc,
        [
            "The final prediction ",
            ("math", subsup(mr("ŷ"), mr("i,t"), mr("(30)"))),
            " is obtained through a linear regression head:",
        ],
    )
    add_equation(doc, eq_yhat())
    add_mixed_para(
        doc,
        [
            "The output is a continuous estimate of the future 30-day excess return of industry ",
            ("math", mr("i")),
            ". This design is more suitable for investment applications than a simple binary classification output because portfolio allocation usually depends on relative expected returns. A continuous prediction can be ranked across industries, transformed into allocation weights, or combined with risk-control rules.",
        ],
    )
    add_mixed_para(
        doc,
        [
            "The training objective is aligned with this ranking-based interpretation. Instead of relying purely on mean squared error, the model uses a hybrid loss that combines pairwise ranking loss ",
            ("math", sub(mr("L"), mr("rank"))),
            " and a weak regression regularization term ",
            ("math", sub(mr("L"), mr("mse"))),
            ". For one graph snapshot, the ranking loss is defined as",
        ],
    )
    add_equation(doc, eq_lrank())
    add_mixed_para(
        doc,
        [
            "This loss penalizes incorrect ordering between industry pairs. If industry ",
            ("math", mr("i")),
            " realizes a higher future excess return than industry ",
            ("math", mr("j")),
            ", the model is encouraged to assign a higher predicted value to industry ",
            ("math", mr("i")),
            ". Thus, the loss directly reflects the cross-sectional ranking objective that is central to sector rotation.",
        ],
    )
    add_equation(doc, eq_lmse())
    add_mixed_para(
        doc,
        [
            "The total loss ",
            ("math", mr("L")),
            " used in the current implementation is",
        ],
    )
    add_equation(doc, eq_ltotal())
    add_mixed_para(
        doc,
        [
            "The relatively large weight on ",
            ("math", sub(mr("L"), mr("rank"))),
            " reflects the investment objective of identifying relative sector strength, while the small weight on ",
            ("math", sub(mr("L"), mr("mse"))),
            " helps stabilize the numerical scale of the predicted excess returns. The parameters of the model, including the linear projection matrices, attention vectors, biases, and output regression weights, are optimized using Adam with early stopping based on validation performance.",
        ],
    )
    add_mixed_para(
        doc,
        [
            "Overall, the DMRGAT algorithm can be summarized as a dynamic graph-based representation learning procedure. For each trading date ",
            ("math", mr("t")),
            ", the model constructs a multi-relation graph ",
            ("math", sub(mr("G"), mr("t"))),
            ", projects industry features into a hidden space, computes edge-weighted attention scores, aggregates neighboring information through multi-head graph attention, applies a second graph attention layer for higher-order integration, and finally predicts future excess returns through a regression head. This architecture is consistent with the financial intuition that sector performance depends not only on each industry's own characteristics but also on the evolving relational structure among industries.",
        ],
    )

    doc.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
