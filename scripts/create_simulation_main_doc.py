from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
CONFIG_PATH = ROOT / "configs" / "default.json"
METRICS_PATH = ROOT / "data" / "processed" / "metrics.json"
ATTENTION_META_PATH = ROOT / "data" / "processed" / "attention_heatmap_meta.json"
OUTPUT_PATH = ROOT / "DMRGAT_Simulation_Main_Section.docx"


def load_json(path: Path) -> dict:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def set_style(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)


def add_para(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)


def fmt(value: float | int | None, digits: int = 4) -> str:
    if value is None:
        return "N/A"
    if isinstance(value, int):
        return str(value)
    return f"{float(value):.{digits}f}"


def add_setting_table(document: Document, config: dict, metrics: dict) -> None:
    rows = [
        ("Industry universe", f"{metrics.get('num_nodes', 15)} core Shenwan level-1 industries"),
        ("Benchmark", "CSI 300 index"),
        ("Forecast horizon", f"{config.get('horizon', 30)} trading days"),
        ("Rolling correlation window", f"{config.get('corr_window', 20)} trading days"),
        ("Graph snapshots", str(metrics.get("num_graphs", "N/A"))),
        ("Train/validation/test split", f"{config.get('train_ratio', 0.8):.2f}/{config.get('val_ratio', 0.1):.2f}/{1 - config.get('train_ratio', 0.8) - config.get('val_ratio', 0.1):.2f}"),
        ("Fusion coefficients", f"alpha={config.get('alpha', 0.45):.2f}, beta={config.get('beta', 0.25):.2f}, gamma={config.get('gamma', 0.30):.2f}"),
        ("Top-k neighbors", str(config.get("top_k_neighbors", 8))),
        ("Hidden dimension / heads", f"{config.get('hidden_dim', 32)} / {config.get('num_heads', 4)}"),
        ("Loss weights", f"ranking={config.get('ranking_loss_weight', 4.0):.2f}, MSE={config.get('mse_loss_weight', 0.05):.2f}"),
    ]
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Specification"
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)

    for item, value in rows:
        cells = table.add_row().cells
        cells[0].text = item
        cells[1].text = value
        for cell in cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)


def add_metric_table(document: Document, metrics: dict) -> None:
    rows = [
        ("Training", metrics.get("train_rmse"), metrics.get("train_mae"), metrics.get("train_ic"), metrics.get("train_rank_ic")),
        ("Validation", metrics.get("val_rmse"), metrics.get("val_mae"), metrics.get("val_ic"), metrics.get("val_rank_ic")),
        ("Test", metrics.get("test_rmse"), metrics.get("test_mae"), metrics.get("test_ic"), metrics.get("test_rank_ic")),
    ]
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Sample", "RMSE", "MAE", "IC", "RankIC"]
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)

    for sample, rmse, mae, ic, rank_ic in rows:
        cells = table.add_row().cells
        values = [sample, fmt(rmse), fmt(mae), fmt(ic), fmt(rank_ic)]
        for cell, value in zip(cells, values):
            cell.text = value
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)


def main() -> None:
    config = load_json(CONFIG_PATH)
    metrics = load_json(METRICS_PATH)
    attention_meta = load_json(ATTENTION_META_PATH)

    doc = Document()
    set_style(doc)
    doc.add_heading("5. Simulation and Empirical Results", level=1)

    add_para(
        doc,
        "This section reports the empirical simulation results of the proposed DMRGAT framework. The purpose of the simulation is to examine whether the dynamic multi-relation graph attention model can produce an informative cross-sectional signal for industry rotation. The experiment is designed as an out-of-sample forecasting task rather than as a trading backtest. Therefore, the evaluation focuses on prediction accuracy and ranking quality, while portfolio construction, transaction costs, and position constraints are left to the application discussion.",
    )
    add_para(
        doc,
        "The simulation follows a chronological research design. All feature construction, target generation, dynamic graph construction, and sample splitting are performed in time order to avoid look-ahead bias. At each valid trading date, the model observes the current industry-level feature matrix and the corresponding fused graph, and predicts the future excess return of each industry over the CSI 300 benchmark during the next forecasting horizon. The final dataset contains a sequence of graph snapshots, each representing one cross section of the industry market.",
    )

    doc.add_heading("5.1 Experimental Setting", level=2)
    add_para(
        doc,
        "The empirical sample is based on daily Chinese market data. The industry universe consists of 15 core Shenwan level-1 industries, selected to represent a compact but economically meaningful cross section of the A-share sector space. The benchmark is the CSI 300 index, and the forecasting target is the future 30-trading-day excess return of each industry relative to the benchmark. The model uses a 20-trading-day rolling window to construct the return-correlation graph and the capital-flow proxy graph. The static industry-prior graph is combined with these two dynamic graphs to form the final adjacency matrix.",
    )
    add_para(
        doc,
        "Table 1 summarizes the main simulation settings. The graph fusion coefficients are kept consistent with the implemented model. The chronological split allocates 80 percent of the valid graph snapshots to training, 10 percent to validation, and the remaining 10 percent to testing. Validation performance is used for early stopping and model selection, while the test set is reserved for final out-of-sample evaluation.",
    )
    doc.add_heading("Table 1. Simulation Settings", level=3)
    add_setting_table(doc, config, metrics)

    doc.add_heading("5.2 Evaluation Metrics", level=2)
    add_para(
        doc,
        "Four metrics are used to evaluate forecasting performance: root mean squared error (RMSE), mean absolute error (MAE), information coefficient (IC), and rank information coefficient (RankIC). RMSE and MAE measure the numerical distance between predicted and realized future excess returns. These metrics are useful for assessing whether the model can estimate the level of future relative performance. However, in sector-rotation applications, point prediction error is not the only relevant criterion.",
    )
    add_para(
        doc,
        "IC is computed as the cross-sectional correlation between predicted excess returns and realized excess returns for each trading date, averaged across dates. RankIC applies the same idea to cross-sectional ranks rather than raw values. RankIC is especially important in this study because the investment task is fundamentally ranking-based. A positive RankIC means that the model tends to place future stronger industries above future weaker industries in the cross section. Therefore, even if numerical prediction errors remain nontrivial, a positive RankIC may still indicate useful sector-selection information.",
    )
    add_para(
        doc,
        "The model-selection criterion is aligned with this interpretation. The validation RankIC is used as the primary early-stopping signal when available. This design is consistent with the ranking-oriented training objective and avoids selecting a model solely because it minimizes in-sample numerical error. Such a choice is reasonable in financial forecasting, where return levels are noisy and unstable, but relative ordering may still carry investment value.",
    )

    doc.add_heading("5.3 Main Forecasting Results of DMRGAT", level=2)
    add_para(
        doc,
        "Table 2 reports the forecasting results of the DMRGAT model on the training, validation, and test samples. The metrics show that the model achieves similar RMSE values on the training and validation sets, while the test RMSE is moderately higher. This pattern suggests that the model does not simply memorize the training sample, but the final out-of-sample period remains more difficult to predict. This is expected in financial data, where market regimes change and industry relationships are not fully stable through time.",
    )
    doc.add_heading("Table 2. DMRGAT Forecasting Performance", level=3)
    add_metric_table(doc, metrics)
    add_para(
        doc,
        f"The validation RankIC is {fmt(metrics.get('val_rank_ic'))}, which indicates that the model learns a positive ranking signal during model selection. The test RankIC is {fmt(metrics.get('test_rank_ic'))}. Although this value is small, it remains positive, suggesting that the model preserves a weak out-of-sample ordering signal. The test IC is {fmt(metrics.get('test_ic'))}, which is also close to zero but positive. These results should be interpreted cautiously: the current model does not produce a strong predictive signal, but it does provide a directionally meaningful ranking tendency under a difficult financial forecasting task.",
    )
    add_para(
        doc,
        "The difference between ranking metrics and point-error metrics is important. The model is trained with a hybrid loss dominated by pairwise ranking loss, so its primary purpose is not to minimize absolute forecast error. Instead, it is designed to rank industries by expected relative strength. A modest positive RankIC is therefore more relevant to the sector-rotation objective than classification accuracy or raw return-level precision. At the same time, the weak magnitude of the signal indicates that further feature engineering, regime conditioning, or temporal graph extensions would be necessary before making stronger economic claims.",
    )

    doc.add_heading("5.4 Attention-Based Interpretation", level=2)
    add_para(
        doc,
        "An advantage of the DMRGAT framework is that the attention mechanism provides an interpretable view of information transmission across industries. After model training, the implementation exports attention heatmaps from the first graph attention layer. These heatmaps summarize how much attention each industry assigns to other industries when updating its node representation. Because the first layer uses multiple attention heads, each head can capture a different relational pattern.",
    )
    add_para(
        doc,
        f"In the current output, the attention visualization is generated from the {attention_meta.get('source_split', 'test')} split on {attention_meta.get('trade_date', 'a selected trading date')}. The layer is {attention_meta.get('layer', 'layer1')}, and the number of attention heads is {attention_meta.get('num_heads', 4)}. The exported files include an average attention heatmap and separate heatmaps for each head. The head-specific heatmaps are useful because one head may emphasize cyclical co-movement, another may focus on technology-related spillovers, while another may capture financial or consumption-related linkages.",
    )
    add_para(
        doc,
        "The attention heatmap should be interpreted as a model-based relational diagnostic rather than as a causal explanation. A high attention value from one industry to another means that, conditional on the learned representation and the fused graph prior, the model relies more heavily on that neighbor's information when updating the current industry node. It does not prove that one industry causes the other to move. Nevertheless, the visualization improves transparency because it allows the researcher to inspect whether the model's learned relations are broadly consistent with economic intuition.",
    )
    add_para(
        doc,
        "This interpretability feature is particularly relevant for journal presentation. In many financial machine-learning applications, predictive models are criticized for being black boxes. The DMRGAT framework partially addresses this concern by linking attention weights to a fused graph constructed from return correlation, fund-flow similarity, and industry-prior relations. As a result, the researcher can examine not only final forecasting metrics but also the relational structure used by the model during prediction.",
    )

    doc.add_heading("5.5 Simulation Summary", level=2)
    add_para(
        doc,
        "The main simulation results suggest that DMRGAT can generate a weak but positive out-of-sample ranking signal for industry rotation. The predictive strength is not large, which is consistent with the low signal-to-noise ratio of financial excess returns. However, the framework provides a coherent empirical pipeline: it constructs dynamic industry graphs, learns edge-weighted attention, optimizes a ranking-oriented objective, and produces interpretable attention heatmaps. These properties make the model useful as a research framework even though the current empirical signal remains modest.",
    )
    add_para(
        doc,
        "The model-comparison subsection can be inserted after this summary or immediately after the main forecasting results. That comparison evaluates DMRGAT against CNN and LSTM baselines under the same target and sample split. The present subsection focuses only on the standalone DMRGAT simulation results and their interpretation, while the separate comparison section discusses whether graph-based modeling provides additional value relative to non-graph neural architectures.",
    )

    doc.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
