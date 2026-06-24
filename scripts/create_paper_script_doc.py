from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "DMRGAT_Paper_Script_English.docx"


TITLE = "Dynamic Multi-Relation Graph Attention Networks for Sector Rotation Forecasting"


SECTIONS = [
    (
        "Introduction",
        [
            "Sector rotation forecasting is a central problem in quantitative asset allocation because investors rarely make decisions on individual assets in isolation. Instead, they evaluate the relative strength of sectors, industries, and styles under changing macroeconomic and market conditions. Traditional forecasting models usually rely on independent time-series signals for each sector and therefore ignore the evolving dependency structure among industries.",
            "This paper develops a Dynamic Multi-Relation Graph Attention Network, or DMRGAT, to predict future sector performance in a relational way. Each industry is treated as a node, while several economically meaningful links are introduced to describe dynamic comovement, capital-flow similarity, and structural industrial relatedness. The key idea is that sector rotation is not only driven by each industry's own price and volume information, but also by the cross-sectional transmission of information across the sector network.",
            "Using this framework, the model is able to learn how shocks, sentiment, and relative momentum propagate between industries. The resulting prediction can provide investors with a ranking-based view of future excess return opportunities rather than a purely univariate forecast. This makes the model suitable for sector allocation, top-down portfolio construction, and broad market style analysis.",
        ],
    ),
    (
        "Basic Model",
        [
            "Let the dynamic sector graph at time t be denoted by G_t = (V, E_t, A_t), where V = {v_1, v_2, ..., v_N} is the set of industry nodes, E_t is the edge set, and A_t is the fused adjacency matrix. In the current implementation, N = 15 because the model is trained on 15 core Shenwan level-1 industries.",
            "For each industry i, the node feature vector is constructed exactly as",
            "x_i,t = [ret_1d_i,t, excess_ret_1d_i,t, ret_std_5_i,t, ret_std_20_i,t, amount_chg_i,t, amount_z20_i,t, mom_5_i,t, mom_20_i,t, ma_gap_i,t, ret_rank_pct_i,t, excess_rank_pct_i,t, mom_rank_pct_5_i,t, mom_rank_pct_20_i,t, amount_rank_pct_i,t, vol_rank_pct_i,t].",
            "These 15 features correspond to the code implementation and include daily return, benchmark-relative excess return, short- and medium-horizon volatility, trading-amount dynamics, momentum, moving-average deviation, and cross-sectional percentile rankings.",
            "The forecasting target is the future 30-day excess return of industry i over the benchmark:",
            "y_i,t^(30) = R_i,t->t+30 - R_m,t->t+30,",
            "where R_i,t->t+30 is the future 30-day return of industry i and R_m,t->t+30 is the future 30-day return of the benchmark index. In the code, this quantity is stored as future_excess_ret and used directly as the regression target.",
            "To describe inter-industry dependence, three relation matrices are introduced. The return-correlation graph is defined by",
            "A_corr,t(i,j) = Corr(ret_1d_i,t-w:t, ret_1d_j,t-w:t),",
            "which captures rolling co-movement in industry returns. The capital-flow similarity graph is defined by",
            "A_fund,t(i,j) = Corr(amount_chg_i,t-w:t, amount_chg_j,t-w:t),",
            "which uses the rolling correlation of trading-amount changes as a proxy for synchronized capital flow. In addition, a static industry prior graph is defined as",
            "A_industry(i,j) = 1, if industries i and j are economically related; 0, otherwise.",
            "The final multi-relation adjacency matrix is constructed as",
            "A_t = alpha * A_corr,t + beta * A_industry + gamma * A_fund,t,",
            "where alpha = 0.45, beta = 0.25, and gamma = 0.30 in the current implementation. Since these coefficients sum to one, the fused graph is a convex combination of the three relation matrices.",
            "Given node features X_t and the fused adjacency A_t, the learning task is to estimate a parametric function",
            "hat(y)_i,t^(30) = f(G_t, X_t; theta),",
            "which predicts the future 30-day excess return for each industry. This defines the precise mathematical learning problem used in the code.",
        ],
    ),
    (
        "Algorithm",
        [
            "Given node features X_t and adjacency A_t, the model starts from h_i^(0) = x_i,t and applies a linear projection to the hidden space:",
            "z_i = W h_i^(0).",
            "For each edge (i, j), the unnormalized graph attention score is computed as",
            "e_ij,t = A_ij,t * LeakyReLU(a^T [z_i || z_j]),",
            "where A_ij,t is the fused edge weight and [z_i || z_j] denotes feature concatenation. This specification is exactly consistent with the implemented weighted attention mechanism.",
            "The normalized attention coefficient is",
            "alpha_ij,t = exp(e_ij,t) / sum_{k in N_i} exp(e_ik,t).",
            "Using the normalized attention weights, the node representation is updated by",
            "h_i^(1) = sigma(sum_{j in N_i} alpha_ij,t z_j).",
            "The first graph attention layer uses M = 4 heads in the current implementation. Therefore, the multi-head output can be written as",
            "h_i^(1) = ||_{m=1}^4 sigma(sum_{j in N_i} alpha_ij,t^(m) z_j^(m)).",
            "The concatenated first-layer output is passed through a second graph attention layer with one head, followed by a linear regression layer to produce the final prediction:",
            "hat(y)_i,t^(30) = W_o h_i^(2).",
            "The training objective combines pairwise ranking loss and mean squared error regularization. For one graph snapshot, the ranking loss is",
            "L_rank = mean_{i,j} log(1 + exp(-(y_i - y_j)(hat(y)_i - hat(y)_j))).",
            "The MSE term is",
            "L_mse = 1/N sum_i (hat(y)_i - y_i)^2.",
            "The final loss used in the code is",
            "L = 4.0 * L_rank + 0.05 * L_mse,",
            "which emphasizes cross-sectional ordering while keeping the predicted return scale numerically stable.",
        ],
    ),
    (
        "Theoretical Analysis",
        [
            "Theoretical intuition for DMRGAT comes from both graph representation learning and financial market structure. When alpha_ij,t is interpreted as an adaptive transmission coefficient, the model can be viewed as learning a time-varying diffusion operator over the sector network. Instead of assuming fixed pairwise influence, DMRGAT learns state-dependent propagation intensity conditioned on both node features and edge priors.",
            "Because the target is future excess return rather than a binary label, relative ordering matters more than exact point accuracy in many allocation problems. If the ranking objective is written as",
            "L_rank = mean_{i,j} log(1 + exp(-(y_i - y_j)(hat(y)_i - hat(y)_j))),",
            "then minimizing L_rank encourages sign(hat(y)_i - hat(y)_j) to be consistent with sign(y_i - y_j). Thus, the optimization procedure directly pushes the model to place future winners above future losers in the cross section.",
            "This property naturally connects the model with RankIC evaluation. If the model output preserves the ordering of future excess returns, then the cross-sectional rank correlation between hat(y)_i and y_i should be positive. Therefore, RankIC is not merely an empirical metric but also a direct reflection of the ranking objective optimized during training.",
            "The multi-relation graph is also theoretically meaningful. Return correlation A_corr,t reflects short-term co-movement, flow similarity A_fund,t reflects synchronized capital behavior, and the prior graph A_industry reflects longer-term economic linkage. Their convex combination",
            "A_t = 0.45 * A_corr,t + 0.25 * A_industry + 0.30 * A_fund,t",
            "acts as a structured inductive bias, restricting attention allocation toward economically plausible neighbors. This can reduce noise relative to an unconstrained attention mechanism and improves interpretability of learned industry interactions.",
        ],
    ),
    (
        "Simulation",
        [
            "The simulation workflow begins with data acquisition, feature engineering, graph construction, model training, and out-of-sample evaluation. Industry-level daily data are collected and aligned with benchmark market data. For each trading date, a dynamic graph is created based on rolling windows, while each node receives both time-series indicators and cross-sectional ranking features.",
            "The dataset is then split chronologically into training, validation, and test sets, ensuring that no future information leaks into earlier periods. During training, the model parameters are optimized by Adam, and the validation set is used for early stopping. This design reflects realistic financial forecasting, where the model is always estimated on the past and evaluated on the future.",
            "Simulation results should be judged using multiple metrics. RMSE and MAE measure numerical forecasting error, while IC and RankIC measure directional and ranking consistency across industries. In sector rotation applications, RankIC is particularly important because it directly reflects whether the model can rank future winners above future losers.",
            "Visual analysis can further strengthen the simulation section. Attention heatmaps from the first graph attention layer reveal which industries the model views as influential neighbors under a given market state. Different attention heads may focus on distinct relational patterns, such as cyclical transmission, technology spillovers, or defensive clustering.",
        ],
    ),
    (
        "Application",
        [
            "In practical investment use, the model's output can be transformed into a ranking signal for industry allocation. At each decision date, sectors with higher predicted future excess returns can be overweighted, while sectors with lower predicted values can be underweighted or excluded. This makes the model suitable for top-down rotation strategies, sector ETF allocation, and institutional tactical asset allocation.",
            "A straightforward application rule is to sort all candidate sectors by predicted future excess return and select the top K sectors for portfolio formation. Position weights can be equal-weighted, score-weighted, or volatility-adjusted. Risk control can be added by limiting exposure to highly correlated industries or by imposing turnover constraints to reduce transaction costs.",
            "Another useful application is interpretability. The attention maps provide a network-level explanation of how the model processes market structure. For example, if banks, non-bank financials, and transportation receive strong mutual attention in a particular period, this may indicate that the model has detected a common macro-sensitive regime. Such insights are valuable for both systematic portfolio managers and discretionary analysts.",
            "Even when predictive performance is moderate, the framework remains useful as a decision-support tool. It can serve as an additional signal layer within a broader investment process that also includes macro judgment, valuation analysis, and risk budgeting. Therefore, the model is not necessarily a stand-alone trading engine, but a structured quantitative component for sector allocation research.",
        ],
    ),
    (
        "Conclusion",
        [
            "This paper proposes a DMRGAT framework for sector rotation forecasting by integrating dynamic return correlation, capital-flow similarity, and industrial structure into a unified graph attention model. Compared with isolated time-series methods, the proposed approach captures cross-sectional information transmission and evolving dependencies among industries.",
            "The model is especially appealing because it aligns well with the logic of investment practice. Sector allocation depends on relative opportunity sets, and graph attention offers a flexible way to translate market structure into predictive signals. By extending the objective from classification to excess-return regression and ranking optimization, the framework becomes more directly relevant for portfolio construction.",
            "Future research may improve the framework in several directions. First, one can add temporal modules such as GRU, LSTM, or temporal graph networks to better capture intertemporal dynamics. Second, macro variables, style factors, valuation metrics, and institutional flow data may be incorporated as richer node or graph features. Third, rolling stability tests and transaction-cost-aware backtests can be added to evaluate whether the learned ranking signal has persistent economic value.",
            "Overall, DMRGAT provides a coherent methodology for combining graph learning and financial sector rotation. Even if raw predictive strength varies across market regimes, the framework offers a meaningful bridge between relational machine learning and practical asset allocation.",
        ],
    ),
]


def set_default_style(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)


def add_title(document: Document, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)


def add_section(document: Document, heading: str, paragraphs: list[str]) -> None:
    document.add_heading(heading, level=1)
    for text in paragraphs:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        run = paragraph.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)


def main() -> None:
    document = Document()
    set_default_style(document)
    add_title(document, TITLE)
    document.add_paragraph()

    for heading, paragraphs in SECTIONS:
        add_section(document, heading, paragraphs)

    output_path = OUTPUT_PATH
    try:
        document.save(output_path)
    except PermissionError:
        output_path = ROOT / "DMRGAT_Paper_Script_English_v2.docx"
        document.save(output_path)
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    main()
