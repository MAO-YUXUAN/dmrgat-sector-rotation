from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "DMRGAT_Section_3_2_Algorithm_Formatted.docx"


def set_default_style(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def add_para(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)


def add_linear_equation(document: Document, latex_like: str) -> None:
    """Insert a centered Office Math paragraph in Word's linear math format."""
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)

    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = latex_like
    run.append(text)
    omath.append(run)
    omath_para.append(omath)
    p._p.append(omath_para)


def main() -> None:
    doc = Document()
    set_default_style(doc)

    add_heading(doc, "3.2 DMRGAT Architecture and Algorithm", level=1)

    add_para(
        doc,
        "The proposed DMRGAT model is designed to transform the dynamic multi-relation industry graph into a set of predictive node representations. Unlike a conventional time-series model that treats each industry independently, DMRGAT explicitly incorporates cross-industry dependence through graph attention. At each trading date t, the model takes the node feature matrix X_t and the fused adjacency matrix A_t as inputs. Each row of X_t corresponds to one industry, and each element of A_t represents the strength of the relation between two industries after integrating return correlation, capital-flow similarity, and static industry linkage. The purpose of the model is to learn a nonlinear mapping from the dynamic graph state (G_t, X_t) to the predicted future excess return for each industry node.",
    )
    add_para(
        doc,
        "The initial node representation of industry i is defined as follows. This notation indicates that the input of the first graph attention layer is exactly the feature vector constructed in the data preprocessing stage.",
    )
    add_linear_equation(doc, "h_i^(0)=x_(i,t)")

    add_para(
        doc,
        "Before computing attention scores, the raw node representation is projected into a hidden feature space through a trainable linear transformation:",
    )
    add_linear_equation(doc, "z_i=W h_i^(0)")

    add_para(
        doc,
        "This projection maps heterogeneous financial variables, including returns, volatility, momentum, liquidity-related information, and cross-sectional rankings, into a common latent representation. It also allows the model to learn task-specific combinations of these variables rather than relying directly on their raw scales. In the current implementation, the hidden dimension is set to 32, which is appropriate for a relatively small industry graph with 15 nodes.",
    )
    add_para(
        doc,
        "For each edge (i, j), the weighted attention score is computed by combining neural attention with the prior graph relation:",
    )
    add_linear_equation(doc, "e_(ij,t)=A_(ij,t)*LeakyReLU(a^T [z_i || z_j])")

    add_para(
        doc,
        "This equation is the key difference between DMRGAT and a vanilla GAT. In a standard GAT, the attention score is determined only by the transformed node features. In DMRGAT, the score is further multiplied by the fused edge weight A_ij,t. Therefore, the model does not allocate attention in a completely unconstrained manner. Instead, attention allocation is guided by economically meaningful prior relations. Industries with stronger return co-movement, capital-flow similarity, or structural linkage receive greater prior influence in the attention mechanism.",
    )
    add_para(
        doc,
        "The attention score is then normalized over the neighborhood of node i through a softmax operation:",
    )
    add_linear_equation(doc, "alpha_(ij,t)=exp(e_(ij,t))/(sum_(k in N_i) exp(e_(ik,t)))")

    add_para(
        doc,
        "The normalized coefficient alpha_ij,t can be interpreted as the relative importance of industry j when updating the representation of industry i. Since the coefficients are normalized within the neighborhood, the attention mechanism produces a probability-like distribution over neighboring industries. This makes the model adaptive to local graph structure and allows different industries to depend on different sets of influential neighbors.",
    )
    add_para(
        doc,
        "Using the normalized attention weights, the first-layer node representation is computed through neighborhood aggregation. In the current implementation, the first graph attention layer uses four attention heads:",
    )
    add_linear_equation(doc, "h_i^(1)=Concat_(m=1)^4 [sigma(sum_(j in N_i) alpha_(ij,t)^(m) z_j^(m))]")

    add_para(
        doc,
        "The multi-head design is meaningful in the sector-rotation setting because industry relationships are heterogeneous. One head may learn cyclical co-movement, another may capture technology-sector spillovers, while another may focus on financial or consumption-related linkages. Instead of forcing all relations into a single attention distribution, the model allows different heads to represent different relational perspectives. The head-specific attention heatmaps generated by the implementation provide an interpretable view of these different relational patterns.",
    )
    add_para(
        doc,
        "After the first graph attention layer, the concatenated representation is normalized, passed through a nonlinear activation, and regularized by dropout. The resulting representation is then passed into a second graph attention layer with one attention head. Conceptually, the first layer aggregates information from directly connected neighboring industries, while the second layer further integrates the information already propagated through the first layer. Therefore, a two-layer GAT structure allows each industry node to incorporate both first-order neighborhood information and higher-order relational signals transmitted through neighboring industries.",
    )
    add_linear_equation(doc, "h_i^(2)=sigma(sum_(j in N_i) alpha_(ij,t)^(2) W^(2) h_j^(1))")

    add_para(
        doc,
        "Here, the superscript in h_i^(1) and h_i^(2) denotes the graph attention layer index rather than the number of training iterations. This distinction is important because the model depth describes representation transformations during one forward pass, while optimization iterations refer to repeated parameter updates during training.",
    )
    add_para(doc, "The final prediction is obtained through a linear regression head:")
    add_linear_equation(doc, "hat(y)_(i,t)^(30)=W_o h_i^(2)")

    add_para(
        doc,
        "The output is a continuous estimate of the future 30-day excess return of industry i. This design is more suitable for investment applications than a simple binary classification output because portfolio allocation usually depends on relative expected returns. A continuous prediction can be ranked across industries, transformed into allocation weights, or combined with risk-control rules.",
    )
    add_para(
        doc,
        "The training objective is aligned with this ranking-based interpretation. Instead of relying purely on mean squared error, the model uses a hybrid loss that combines pairwise ranking loss and a weak regression regularization term. For one graph snapshot, the ranking loss is defined as",
    )
    add_linear_equation(doc, "L_rank=mean_(i,j) log(1+exp(-(y_i-y_j)(hat(y)_i-hat(y)_j)))")

    add_para(
        doc,
        "This loss penalizes incorrect ordering between industry pairs. If industry i realizes a higher future excess return than industry j, the model is encouraged to assign a higher predicted value to industry i. Thus, the loss directly reflects the cross-sectional ranking objective that is central to sector rotation.",
    )
    add_linear_equation(doc, "L_mse=1/N sum_i (hat(y)_i-y_i)^2")

    add_para(doc, "The total loss used in the current implementation is")
    add_linear_equation(doc, "L=4.0 L_rank+0.05 L_mse")

    add_para(
        doc,
        "The relatively large weight on the ranking loss reflects the investment objective of identifying relative sector strength, while the small MSE term helps stabilize the numerical scale of the predicted excess returns. The parameters of the model, including the linear projection matrices, attention vectors, biases, and output regression weights, are optimized using Adam with early stopping based on validation performance.",
    )
    add_para(
        doc,
        "Overall, the DMRGAT algorithm can be summarized as a dynamic graph-based representation learning procedure. For each trading date, the model constructs a multi-relation graph, projects industry features into a hidden space, computes edge-weighted attention scores, aggregates neighboring information through multi-head graph attention, applies a second graph attention layer for higher-order integration, and finally predicts future excess returns through a regression head. This architecture is consistent with the financial intuition that sector performance depends not only on each industry's own characteristics but also on the evolving relational structure among industries.",
    )

    doc.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
