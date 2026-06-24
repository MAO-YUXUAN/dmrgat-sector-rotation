from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
CONFIG_PATH = ROOT / "configs" / "default.json"
METRICS_PATH = ROOT / "data" / "processed" / "metrics.json"
OUTPUT_PATH = ROOT / "DMRGAT_Application_Section.docx"


def load_json(path: Path) -> dict:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def set_style(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)


def add_para(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)


def add_bullet(document: Document, text: str) -> None:
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)


def fmt(value: float | int | None, digits: int = 4) -> str:
    if value is None:
        return "N/A"
    if isinstance(value, int):
        return str(value)
    return f"{float(value):.{digits}f}"


def add_signal_table(document: Document) -> None:
    rows = [
        ("Step 1", "Generate predictions", "Use the trained DMRGAT model to estimate future 30-day excess return for each industry."),
        ("Step 2", "Rank industries", "Sort industries by predicted excess return from high to low."),
        ("Step 3", "Select candidates", "Focus on top-ranked industries as overweight candidates and bottom-ranked industries as underweight candidates."),
        ("Step 4", "Apply risk filters", "Check volatility, concentration, liquidity, and industry exposure constraints before allocation."),
        ("Step 5", "Monitor signal stability", "Update signals regularly and avoid excessive turnover caused by small ranking changes."),
    ]
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["Stage", "Operation", "Investment Interpretation"]
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)

    for stage, operation, interpretation in rows:
        cells = table.add_row().cells
        values = [stage, operation, interpretation]
        for cell, value in zip(cells, values):
            cell.text = value
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)


def main() -> None:
    config = load_json(CONFIG_PATH)
    metrics = load_json(METRICS_PATH)
    horizon = int(config.get("horizon", 30))
    num_nodes = int(metrics.get("num_nodes", 15))
    rank_ic = metrics.get("test_rank_ic")
    test_rmse = metrics.get("test_rmse")

    doc = Document()
    set_style(doc)
    doc.add_heading("6. Application: Sector-Rotation Decision Support", level=1)

    add_para(
        doc,
        "The proposed DMRGAT framework is designed not only as a forecasting model but also as a decision-support tool for sector rotation. In practical asset allocation, investors often need to determine which industries should be overweighted, which should be kept neutral, and which should be underweighted. This decision is inherently cross-sectional. Therefore, the most relevant output of the model is not merely the absolute predicted value of future returns, but the relative ranking of industries according to their expected excess performance over the benchmark.",
    )
    add_para(
        doc,
        f"In the current implementation, the model predicts each industry's future {horizon}-trading-day excess return relative to the CSI 300 benchmark. The industry universe contains {num_nodes} core Shenwan level-1 industries. For each valid trading date, the trained model produces one predicted excess-return score for each industry. These scores can be sorted cross-sectionally to form an industry-strength ranking. The top-ranked sectors may be interpreted as potential overweight candidates, while the bottom-ranked sectors may be treated as underweight or avoidance candidates.",
    )

    doc.add_heading("6.1 From Prediction Scores to Ranking Signals", level=2)
    add_para(
        doc,
        "A direct application of DMRGAT is the construction of a ranking signal. Let the model output a vector of predicted excess returns for all industries at date t. The industries are sorted from the highest predicted value to the lowest predicted value. This ranking can be used as a relative-strength indicator. Compared with a binary classification signal, a continuous predicted excess-return score is more flexible because it preserves the intensity of the model's view. Two industries may both be expected to outperform the benchmark, but the one with the higher predicted excess return should receive stronger consideration in allocation decisions.",
    )
    add_para(
        doc,
        "The ranking signal can be used in several practical ways. A simple approach is to select the top-K industries as candidate overweight sectors. Another approach is to divide industries into quantile groups, such as top, middle, and bottom groups, and then assign portfolio tilts according to their relative scores. A more conservative approach is to use the DMRGAT ranking only as a screening tool, combining it with macroeconomic views, valuation indicators, or risk-budget constraints before making final allocation decisions.",
    )
    doc.add_heading("Table 1. Practical Signal-Generation Procedure", level=3)
    add_signal_table(doc)

    doc.add_heading("6.2 Portfolio Construction Framework", level=2)
    add_para(
        doc,
        "Although this study does not conduct a formal trading backtest, the model output can be naturally embedded into a portfolio-construction framework. The first step is to transform predicted excess returns into allocation scores. These scores may be standardized cross-sectionally to reduce the effect of unstable return scales. The second step is to convert the standardized scores into portfolio weights. For example, industries with positive standardized scores may receive overweight positions, while those with negative scores may receive underweight positions relative to a benchmark allocation.",
    )
    add_para(
        doc,
        "A practical allocation rule should include several constraints. First, position concentration should be limited so that the portfolio is not dominated by one or two industries. Second, turnover should be controlled because frequent re-ranking may lead to excessive transaction costs. Third, liquidity conditions should be considered, especially when applying the signal to industry ETFs or sector funds. Fourth, risk exposure should be monitored to avoid unintended concentration in macro factors such as commodity prices, interest rates, or technology growth sentiment.",
    )
    add_para(
        doc,
        "For these reasons, the model should be viewed as a signal-generation component rather than a complete trading system. A full investment strategy would require additional modules, including portfolio optimization, transaction-cost modeling, stop-loss or drawdown controls, and benchmark-relative risk management. The present paper focuses on the predictive and ranking capability of the graph-learning model; the translation from prediction to actual portfolio weights is left as an applied extension.",
    )

    doc.add_heading("6.3 Interpretation of Attention for Investment Research", level=2)
    add_para(
        doc,
        "Another application of DMRGAT is interpretability. The attention heatmaps generated by the model provide a structured view of inter-industry information transmission. For each attention head, the heatmap shows how much information one industry receives from other industries during representation updating. This is valuable for investment research because sector rotation is often driven by cross-industry spillovers. For example, movements in upstream cyclical industries may influence manufacturing sectors, while technology-related industries may exhibit strong mutual dependence during growth-oriented market regimes.",
    )
    add_para(
        doc,
        "The attention visualization should not be interpreted as causal evidence. A high attention weight does not prove that one industry causes another industry to outperform. Instead, it indicates that the trained model relies more heavily on that relation when forming predictions. This distinction is important for responsible application. The heatmap is best used as a diagnostic tool: it helps analysts examine whether the model's learned structure is economically plausible, whether attention is overly concentrated, and whether different heads capture different types of industry linkage.",
    )
    add_para(
        doc,
        "In an investment workflow, attention heatmaps can support qualitative review. If the model predicts strong performance for a particular sector, analysts can inspect which neighboring industries contribute to that signal. If the influential neighbors are economically related, the signal may be more credible. If the attention pattern appears unstable or economically unreasonable, the prediction may require additional scrutiny. In this way, DMRGAT offers not only numerical forecasts but also a transparent relational layer for model interpretation.",
    )

    doc.add_heading("6.4 Practical Limitations and Risk Controls", level=2)
    add_para(
        doc,
        f"The empirical results indicate that the current model produces only a weak out-of-sample ranking signal. The test RankIC is {fmt(rank_ic)}, and the test RMSE is {fmt(test_rmse)}. These values suggest that the model should not be used as a standalone trading engine. Instead, it should be treated as one component in a broader decision-making process. This limitation is common in financial forecasting because future excess returns are noisy, regime-dependent, and influenced by macroeconomic shocks that may not be fully captured by historical industry data.",
    )
    add_para(
        doc,
        "Several risk controls are necessary before practical deployment. The first is signal smoothing. Small changes in predicted rankings should not automatically trigger large portfolio adjustments. The second is regime awareness. A model trained on historical relationships may perform differently during crisis periods, policy-driven markets, or liquidity shocks. The third is benchmark-relative exposure control. Since the target is excess return over the CSI 300, the application should focus on relative industry tilts rather than unconstrained absolute bets. The fourth is robustness monitoring. If RankIC deteriorates persistently, the model should be recalibrated or temporarily deactivated.",
    )
    add_bullet(doc, "Use the model primarily as a ranking signal rather than a precise return forecast.")
    add_bullet(doc, "Combine model rankings with macro, valuation, liquidity, and risk-control information.")
    add_bullet(doc, "Avoid excessive turnover by applying ranking thresholds or signal smoothing.")
    add_bullet(doc, "Monitor out-of-sample RankIC regularly to detect signal decay.")
    add_bullet(doc, "Treat attention heatmaps as interpretability diagnostics, not causal explanations.")

    doc.add_heading("6.5 Application Summary", level=2)
    add_para(
        doc,
        "The application value of DMRGAT lies in its ability to transform dynamic industry relationships into a ranking-oriented sector signal. The model integrates node-level market features with inter-industry dependence, produces continuous future excess-return predictions, and offers attention-based interpretation. These properties make it suitable for research-oriented sector allocation, ranking-based screening, and analyst decision support. However, the empirical signal remains modest, so the model should be applied cautiously and combined with additional investment constraints.",
    )
    add_para(
        doc,
        "In summary, DMRGAT provides a bridge between graph representation learning and practical sector-rotation research. Its output can guide industry ranking, support benchmark-relative allocation tilts, and help analysts understand relational patterns among sectors. The model is not a complete portfolio strategy by itself, but it offers a structured and interpretable foundation for future extensions involving transaction-cost-aware backtesting, macro-factor integration, and dynamic risk-controlled allocation.",
    )

    doc.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
