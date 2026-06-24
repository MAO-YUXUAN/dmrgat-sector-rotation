from __future__ import annotations

import json
from html import escape
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
CONFIG_PATH = ROOT / "configs" / "default.json"
METRICS_PATH = ROOT / "data" / "processed" / "metrics.json"
OUTPUT_PATH = ROOT / "DMRGAT_Section_3_1_Basic_Model_OMML_Equations.docx"


def load_json(path: Path) -> dict:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def mr(text: str) -> str:
    return f"<m:r><m:t>{escape(text)}</m:t></m:r>"


def sub(base: str, lower: str) -> str:
    return f"<m:sSub><m:e>{base}</m:e><m:sub>{lower}</m:sub></m:sSub>"


def sup(base: str, upper: str) -> str:
    return f"<m:sSup><m:e>{base}</m:e><m:sup>{upper}</m:sup></m:sSup>"


def subsup(base: str, lower: str, upper: str) -> str:
    return f"<m:sSubSup><m:e>{base}</m:e><m:sub>{lower}</m:sub><m:sup>{upper}</m:sup></m:sSubSup>"


def frac(num: str, den: str) -> str:
    return f"<m:f><m:num>{num}</m:num><m:den>{den}</m:den></m:f>"


def omath(expr: str) -> str:
    return f'<m:oMathPara {nsdecls("m")}><m:oMath>{expr}</m:oMath></m:oMathPara>'


def add_equation(document: Document, expr: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p._p.append(parse_xml(omath(expr)))


def add_para(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)


def add_mixed_para(document: Document, segments: list[str | tuple[str, str]]) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    for segment in segments:
        if isinstance(segment, tuple):
            kind, expr = segment
            if kind != "math":
                raise ValueError(f"Unknown segment kind: {kind}")
            p._p.append(parse_xml(f'<m:oMath {nsdecls("m")}>{expr}</m:oMath>'))
        else:
            r = p.add_run(segment)
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)


def set_style(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)


def feature(name: str) -> str:
    return sub(mr(name), mr("i,t"))


def eq_graph(num_nodes: int) -> str:
    left = sub(mr("G"), mr("t"))
    graph = mr("(") + mr("V") + mr(",") + sub(mr("E"), mr("t")) + mr(",") + sub(mr("A"), mr("t")) + mr(")")
    node_set = (
        mr(",  V={")
        + sub(mr("v"), mr("1"))
        + mr(",")
        + sub(mr("v"), mr("2"))
        + mr(",...,")
        + sub(mr("v"), mr("N"))
        + mr("},  N=")
        + mr(str(num_nodes))
    )
    return left + mr("=") + graph + node_set


def eq_feature_vector() -> str:
    return sub(mr("x"), mr("i,t")) + mr("=") + mr("[")


def eq_feature_row(names: list[str], is_last: bool = False) -> str:
    body = mr("")
    for idx, name in enumerate(names):
        body += feature(name)
        if idx != len(names) - 1:
            body += mr(", ")
    return body + (mr("]") if is_last else mr(","))


def eq_target(horizon: int) -> str:
    y = subsup(mr("y"), mr("i,t"), mr(f"({horizon})"))
    ri = sub(mr("R"), mr(f"i,t:t+{horizon}"))
    rm = sub(mr("R"), mr(f"m,t:t+{horizon}"))
    return y + mr("=") + ri + mr("-") + rm


def eq_corr_raw(window: int) -> str:
    left = sub(mr("C_corr"), mr("t")) + mr("(") + mr("i,j") + mr(")")
    rhs = (
        mr("Corr(")
        + sub(mr("ret_1d"), mr(f"i,t-{window-1}:t"))
        + mr(", ")
        + sub(mr("ret_1d"), mr(f"j,t-{window-1}:t"))
        + mr(")")
    )
    return left + mr("=") + rhs


def eq_fund_raw(window: int) -> str:
    left = sub(mr("C_fund"), mr("t")) + mr("(") + mr("i,j") + mr(")")
    rhs = (
        mr("Corr(")
        + sub(mr("amount_chg"), mr(f"i,t-{window-1}:t"))
        + mr(", ")
        + sub(mr("amount_chg"), mr(f"j,t-{window-1}:t"))
        + mr(")")
    )
    return left + mr("=") + rhs


def eq_corr_map(prefix: str) -> str:
    a = sub(mr(f"A_{prefix}"), mr("t")) + mr("(") + mr("i,j") + mr(")")
    c = sub(mr(f"C_{prefix}"), mr("t")) + mr("(") + mr("i,j") + mr(")")
    return a + mr("=") + frac(c + mr("+1"), mr("2"))


def eq_industry() -> str:
    left = mr("A_industry") + mr("(") + mr("i,j") + mr(")")
    rhs = mr("1 if industries i and j are economically related; 0 otherwise")
    return left + mr("=") + rhs


def eq_fusion(alpha: float, beta: float, gamma: float) -> str:
    return (
        sub(mr("A"), mr("t"))
        + mr("=")
        + mr(f"{alpha:.2f}")
        + sub(mr("A_corr"), mr("t"))
        + mr("+")
        + mr(f"{beta:.2f}")
        + mr("A_industry")
        + mr("+")
        + mr(f"{gamma:.2f}")
        + sub(mr("A_fund"), mr("t"))
    )


def eq_convex(alpha: float, beta: float, gamma: float) -> str:
    return mr(f"{alpha:.2f}+{beta:.2f}+{gamma:.2f}=1")


def main() -> None:
    config = load_json(CONFIG_PATH)
    metrics = load_json(METRICS_PATH)

    horizon = int(config.get("horizon", 30))
    corr_window = int(config.get("corr_window", 20))
    alpha = float(config.get("alpha", 0.45))
    beta = float(config.get("beta", 0.25))
    gamma = float(config.get("gamma", 0.30))
    top_k = int(config.get("top_k_neighbors", 8))
    num_graphs = int(metrics.get("num_graphs", 2346))
    num_nodes = int(metrics.get("num_nodes", len(config.get("core_industries", [])) or 15))
    industries = config.get("core_industries", [])
    industry_text = ", ".join(industries)

    doc = Document()
    set_style(doc)
    doc.add_heading("3.1 Problem Setup and Basic Model", level=1)

    add_mixed_para(
        doc,
        [
            "This section formalizes the sector-rotation forecasting problem before introducing the graph attention architecture. The empirical object is a daily cross section of Shenwan level-1 industries in the Chinese A-share market. Rather than treating each industry as an isolated time-series prediction unit, the proposed framework represents the market as a dynamic graph whose nodes correspond to industries and whose edges describe time-varying dependence. At each trading date ",
            ("math", mr("t")),
            ", the graph is denoted as",
        ],
    )
    add_equation(doc, eq_graph(num_nodes))
    add_mixed_para(
        doc,
        [
            "where ",
            ("math", mr("V")),
            " is the industry node set, ",
            ("math", sub(mr("E"), mr("t"))),
            " is the time-varying edge set, and ",
            ("math", sub(mr("A"), mr("t"))),
            " is the fused weighted adjacency matrix used by the model. In the current implementation, the universe contains ",
            ("math", mr(str(num_nodes))),
            " core industries selected from the Shenwan level-1 classification. These industries are: ",
            industry_text,
            ". The selected universe covers cyclical sectors, manufacturing sectors, technology-oriented sectors, consumption sectors, health care, finance, and transportation. This reduced but economically representative universe is used because sector rotation is fundamentally a cross-sectional allocation problem, and a compact universe helps reduce estimation noise in rolling correlations and graph attention weights.",
        ],
    )
    add_mixed_para(
        doc,
        [
            "For each industry ",
            ("math", mr("i")),
            " and date ",
            ("math", mr("t")),
            ", the node feature vector is constructed directly from the implemented preprocessing pipeline. The vector is",
        ],
    )
    add_equation(doc, eq_feature_vector())
    add_equation(
        doc,
        eq_feature_row(["ret_1d", "excess_ret_1d", "ret_std_5", "ret_std_20", "amount_chg"]),
    )
    add_equation(
        doc,
        eq_feature_row(["amount_z20", "mom_5", "mom_20", "ma_gap", "ret_rank_pct"]),
    )
    add_equation(
        doc,
        eq_feature_row(
            ["excess_rank_pct", "mom_rank_pct_5", "mom_rank_pct_20", "amount_rank_pct", "vol_rank_pct"],
            is_last=True,
        ),
    )
    add_para(
        doc,
        "The feature design combines time-series information and cross-sectional information. The variables ret_1d and excess_ret_1d describe the latest absolute and benchmark-relative performance of an industry. The variables ret_std_5 and ret_std_20 measure short- and medium-window realized volatility, which helps the model distinguish stable trends from noisy price movements. The variables amount_chg and amount_z20 capture changes in trading amount and the standardized intensity of recent trading activity. These variables are used as liquidity and participation proxies because industry-level fund-flow information is not always directly observable in a clean and continuous form. Momentum is represented by mom_5 and mom_20, while ma_gap measures the deviation of the current price level from a moving-average reference. These variables summarize persistence and trend-following information that is commonly used in tactical allocation research."
    )
    add_para(
        doc,
        "The remaining six variables are cross-sectional percentile ranks. Specifically, ret_rank_pct, excess_rank_pct, mom_rank_pct_5, mom_rank_pct_20, amount_rank_pct, and vol_rank_pct describe where an industry stands relative to other industries on the same trading date. These features are important because sector rotation decisions are inherently relative. A sector does not need to have a high absolute return to be attractive; it needs to be stronger than competing sectors after accounting for market conditions and risk. By including both raw time-series features and cross-sectional ranks, the input representation allows the model to learn from each industry's own dynamics as well as from its position within the market-wide industry distribution."
    )
    add_mixed_para(
        doc,
        [
            "The forecasting target is specified as a continuous future excess return rather than a binary up-or-down label. For horizon ",
            ("math", mr(str(horizon))),
            ", the target of industry ",
            ("math", mr("i")),
            " at date ",
            ("math", mr("t")),
            " is defined as",
        ],
    )
    add_equation(doc, eq_target(horizon))
    add_mixed_para(
        doc,
        [
            "where ",
            ("math", sub(mr("R"), mr(f"i,t:t+{horizon}"))),
            " is the future ",
            ("math", mr(str(horizon))),
            "-trading-day return of industry ",
            ("math", mr("i")),
            ", and ",
            ("math", sub(mr("R"), mr(f"m,t:t+{horizon}"))),
            " is the corresponding return of the CSI 300 benchmark. In the code, this target is stored as future_excess_ret and is used directly as the regression target. This choice is more consistent with investment use than the earlier binary classification target because allocation decisions usually depend on relative expected returns. A continuous forecast can be ranked, thresholded, or combined with portfolio constraints, whereas a binary label only indicates whether an industry is expected to outperform the benchmark.",
        ],
    )
    add_para(
        doc,
        "The graph component is designed to encode inter-industry dependence from three complementary perspectives. The first relation is return co-movement. Industries that move together over a recent rolling window may share macro drivers, supply-chain exposure, risk appetite sensitivity, or common institutional positioning. The raw rolling return correlation is defined as"
    )
    add_equation(doc, eq_corr_raw(corr_window))
    add_mixed_para(
        doc,
        [
            "Because raw correlations lie in ",
            ("math", mr("[-1,1]")),
            ", the implementation maps the correlation into a nonnegative edge-strength matrix before graph fusion:",
        ],
    )
    add_equation(doc, eq_corr_map("corr"))
    add_para(
        doc,
        "This transformation preserves the ordering of correlations while making the adjacency entries compatible with the attention mechanism, where edge weights act as multiplicative priors on attention scores. A high value means that two industries have recently exhibited similar return behavior; a low value means that their return dynamics have been weakly related or negatively related during the rolling window."
    )
    add_para(
        doc,
        "The second relation is based on trading-amount dynamics and serves as a capital-flow proxy. Since direct industry-level fund-flow measures may be unavailable or inconsistent across the full sample, the implemented model uses rolling correlation in amount_chg to approximate synchronized changes in market participation. The raw relation is"
    )
    add_equation(doc, eq_fund_raw(corr_window))
    add_equation(doc, eq_corr_map("fund"))
    add_para(
        doc,
        "This fund-flow relation is economically distinct from return correlation. Two industries may not have identical return movements but may still experience simultaneous increases in trading activity, suggesting a common shift in investor attention or liquidity allocation. Including this relation allows the graph to reflect not only price co-movement but also synchronized changes in trading intensity."
    )
    add_para(
        doc,
        "The third relation is a static industry prior graph based on economic linkage. Unlike the first two matrices, which are re-estimated over rolling windows, this matrix reflects relatively stable structural relationships among industries, such as upstream-downstream production links, common demand exposure, technology diffusion, or financial-market linkage. It is specified as a binary matrix:"
    )
    add_equation(doc, eq_industry())
    add_mixed_para(
        doc,
        [
            "The final adjacency matrix is a convex fusion of the three relation channels. With the implemented coefficients ",
            ("math", mr(f"{alpha:.2f}")),
            ", ",
            ("math", mr(f"{beta:.2f}")),
            ", and ",
            ("math", mr(f"{gamma:.2f}")),
            ", the fused graph is",
        ],
    )
    add_equation(doc, eq_fusion(alpha, beta, gamma))
    add_equation(doc, eq_convex(alpha, beta, gamma))
    add_mixed_para(
        doc,
        [
            "The convex condition means that ",
            ("math", sub(mr("A"), mr("t"))),
            " is a weighted average of dynamic return correlation, static industry linkage, and dynamic fund-flow similarity. This is useful both mathematically and economically. Mathematically, it keeps the scale of the adjacency matrix controlled after normalization. Economically, it ensures that each edge weight remains interpretable as a combination of three meaningful relation sources rather than an unrestricted numerical score. In the implementation, only the strongest neighbors are retained for each node, with ",
            ("math", mr(f"k={top_k}")),
            ", which reduces noisy weak edges and makes the graph attention operation more focused.",
        ],
    )
    add_mixed_para(
        doc,
        [
            "After rolling-window construction, feature alignment, target generation, and graph filtering, the final dataset contains ",
            ("math", mr(str(num_graphs))),
            " valid graph snapshots. Each snapshot contains the node feature matrix ",
            ("math", sub(mr("X"), mr("t"))),
            ", the fused adjacency matrix ",
            ("math", sub(mr("A"), mr("t"))),
            ", and the future excess-return target vector ",
            ("math", subsup(mr("y"), mr("t"), mr(f"({horizon})"))),
            ". The learning task is therefore to estimate a parametric mapping from graph-structured market information to future industry relative performance. This setup provides the basic model foundation for the DMRGAT architecture in the following section: the node features describe individual industry states, the fused adjacency matrix describes inter-industry dependence, and the target defines the investment-relevant quantity to be predicted.",
        ],
    )
    add_para(
        doc,
        "In summary, the basic model transforms sector rotation from a set of independent forecasting problems into a dynamic graph prediction problem. This formulation is particularly suitable for financial markets because industry performance is rarely independent. Macroeconomic shocks, policy changes, commodity cycles, credit conditions, and investor risk appetite often propagate across economically related sectors. By combining node-level market features with relation-level graph information, the model is able to represent both the state of each industry and the structure through which information may spread across industries. The subsequent DMRGAT algorithm builds on this setup by learning attention-based aggregation weights over the fused graph and by optimizing predictions under a ranking-oriented objective."
    )

    doc.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
